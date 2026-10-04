"""planilha_envio.py — planilha de envio em .xlsx (Etapa E2b), três abas.
Motivo: o CSV com vírgula do QUALIFICADOR abre numa coluna só no Excel do
diretor — .xlsx real, formatado, resolve isso de vez.

MVP de mensagem por MODELO FIXO (mudança de rumo, decisão do diretor,
24/09/2026): a primeira mensagem de cada lead sai de `angulo_mensagem.py`
(ângulo escolhido por regra + 5 textos fixos), nunca mais de uma chamada de
LLM. A Etapa E2 (`pipeline_estrategia_mensagem.py`, `entrada_estrategia.py`,
`prompts/estrategia_mensagem_sistema.md`) e os templates antigos da pista
Direta (`templates_direta.py`, 5 classes) continuam existindo no código e
nos testes — só saem do fluxo desta planilha.

Aba "Geral": todas as linhas do CSV humano mais recente do QUALIFICADOR
(pistas Direta/Espera) -- ângulo `sem_site`/`doctoralia`.

Aba "Nata": `nata` + `candidatos_triagem` do lote mais recente (coluna
"pista" distingue as duas) -- ângulo `contato`/`reputacao`/`lentidao`.

Aba "Como usar": texto fixo com as regras do operador (não gerado por
lead).

Lead sem ângulo (nenhuma regra bateu): sem mensagem, coluna `Status` =
"sem mensagem gerada" (Nata) ou célula vazia (Geral, como antes). Mensagem
que falha no validador: célula vazia + `Aviso` com o motivo -- nunca uma
mensagem inválida na planilha.

Nunca sobrescreve uma planilha anterior. Exclui, na geração, quem já está no
registro de abordados (`registro_abordagens.py`) — e informa quantos.
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

try:
    import eqc
except ModuleNotFoundError:  # pragma: no cover - bootstrap de sys.path
    import sys as _sys

    _sys.path.insert(0, str(Path(__file__).resolve().parent))
    import eqc

import angulo_mensagem
import contrato_loader
import email_utils
import entrada_estrategia
import env_loader
import registro_abordagens
import whatsapp_utils
from validador_mensagem import carregar_config as carregar_config_validacao

_DADOS_DIR = Path(__file__).resolve().parent.parent / "_planilhas"
DIRETORIO_SAIDA_PADRAO = _DADOS_DIR

# Entrega CONSOLIDADA pro diretor, fora do repo: um único arquivo (nunca
# versionado por timestamp) que cada rodada ABRE e ACRESCENTA -- nunca recria
# do zero (decisão do diretor, 01/10/2026). Só duas abas (Geral, Nata); sem
# "Como usar". Lida só de COMERCIAL_PLANILHAS_DIR (comercial/.env ou env real);
# ausente = cópia pulada, a rodada nunca falha por causa dela -- mesmo padrão
# de QUALIFICADOR_SAIDA_HUMANA_DIR em qualificador/prospeccao_ia/saida_humana.py.
ENV_CONSOLIDADA_DIR = "COMERCIAL_PLANILHAS_DIR"
NOME_PLANILHA_CONSOLIDADA = "planilha_envio_consolidada.xlsx"

_RE_TIMESTAMP = re.compile(r"(\d{8}-\d{6})")
_RE_PROFISSIONAL = re.compile(r"^\s*(?:dr\.?|dra\.?|doctor|doctora)\b", re.IGNORECASE)

COLUNAS_CSV_HUMANO = (
    "pista", "motivo", "nome", "cidade", "telefone", "email", "instagram", "site",
    "classe_site", "avaliacoes", "nota", "google_maps_url", "place_id",
)

# "Respondeu" (livre) virou o funil estruturado (Etapa/Resultado/Mensagem/
# Motivo da edição/Decidi não enviar/Motivo de não enviar) -- decisão do
# diretor, 24/09/2026, quarta rodada: registro_abordagens.xlsx passa a ser a
# planilha de acompanhamento, mesmas colunas nas duas abas e no registro
# (registro_abordagens.COLUNAS_FUNIL é a única fonte -- nunca duplicada
# aqui). "Data 1º contato" e "Próximo contato" são novas (quinta rodada,
# mesmo dia). "Enviado em" saiu (sexta rodada, mesmo dia) -- "Data 1º
# contato" passou a ser o único sinal de envio, lido por
# registro_abordagens.py. "Follow-up em"/"Observação" continuam.
COLUNAS_DIRETOR = (
    ("Data 1º contato", "Próximo contato")
    + registro_abordagens.COLUNAS_FUNIL
    + ("Follow-up em", "Observação")
)

# "place_id" é o identificador padrão de todo lead, em TODA aba de leads
# (decisão do diretor, 25/09/2026, defeito grave: a Nata não tinha essa
# coluna -- registro_abordagens._linhas_enviadas pulava a aba em silêncio,
# e nenhum envio dela seria registrado). Sempre por último, sempre
# preenchida -- gerar_planilha recusa gravar qualquer linha sem ela
# (PlaceIdAusenteError). Oculta na planilha (_escrever_aba) -- é para o
# programa, não para o diretor editar.
#
# "WhatsApp" é o ÚNICO link de WhatsApp por linha (decisão do diretor,
# 25/09/2026, especificação final do MVP): com a mensagem sugerida, quando
# houver; senão, link de conversa vazia (quando há celular válido). A coluna
# "Abrir conversa" saiu -- duplicava o mesmo link.
COLUNAS_GERAL = (
    "pista", "motivo", "nome", "cidade", "telefone", "WhatsApp", "email", "Abrir e-mail", "Canal",
    "instagram", "site", "classe_site", "avaliacoes", "nota", "google_maps_url",
    "Ângulo", "Mensagem sugerida", "Profissional detectado", "Captura",
) + COLUNAS_DIRETOR + ("Aviso", "Contexto para IA", "place_id")

COLUNAS_NATA = (
    "pista", "nome", "telefone", "WhatsApp", "email", "Abrir e-mail", "Canal",
    "problema vendável", "Ângulo", "mensagem", "Status",
    "conferir_antes_de_enviar", "Profissional detectado", "Captura",
    "site", "google_maps_url", "texto_site",
) + COLUNAS_DIRETOR + ("Aviso", "Contexto para IA", "place_id")

STATUS_SEM_MENSAGEM = "sem mensagem gerada"

# Limite real do Excel por célula (.xlsx) -- acima disso o arquivo corrompe a
# célula. "Contexto para IA" é a única coluna que pode chegar perto (o texto
# do site sozinho já vai a até 4000 chars, config/texto_site.json do
# QUALIFICADOR) -- por isso é a única parte truncada quando preciso.
LIMITE_CELULA_XLSX = 32767
_AVISO_TRUNCAMENTO = " [texto do site truncado -- limite de 32.767 caracteres por célula do Excel]"

_COMO_USAR_TEXTO = (
    ("Lentidão", "Antes de enviar, abra o site no seu celular; se carregar rápido, não envie."),
    (
        "Profissional detectado",
        'Quando a coluna "Profissional detectado" = "sim" e você tiver certeza do sobrenome, '
        'troque "Hola, buenas" por "Hola, Dr./Dra. [Apellido]".',
    ),
    (
        "Classificar a resposta",
        "resposta positiva: \"Sí, envíamelo\", \"¿Cuánto cobras?\", \"Sí, cuéntame\", \"Me interesa\", "
        "\"¿Qué cambiarías?\". resposta neutra: \"¿Quién eres?\" (útil, mas não positiva). resposta "
        "negativa: \"No me interesa\", \"No necesito nada\". \"No contactar\" = pediu para não "
        "contatar. Nenhuma resposta = sem resposta.",
    ),
)


class PlanilhaJaExisteError(Exception):
    """Já existe uma planilha com este nome — nunca sobrescrever uma rodada anterior."""


class PlaceIdAusenteError(Exception):
    """Uma ou mais linhas de uma aba de leads sem `place_id` (decisão do
    diretor, 25/09/2026, defeito grave). `place_id` é o identificador padrão
    usado por `registro_abordagens.py` para registrar envios e excluir leads
    já abordados -- uma linha sem ele nunca é gravada; a planilha inteira
    falha, com o nome da aba e as linhas afetadas, em vez de gerar uma
    planilha incompleta em silêncio."""


class LinkInconsistenteError(Exception):
    """Uma ou mais linhas têm o hyperlink (destino do clique) de WhatsApp ou
    e-mail desalinhado dos dados da própria linha -- defeito real, achado
    01/10/2026: `ws.delete_rows()` (usado numa limpeza manual da planilha
    consolidada) desloca o TEXTO das células mas não realinha o objeto
    `hyperlink` do openpyxl, que fica apontando pra posição anterior. Como o
    operador clica direto no link (WhatsApp/e-mail abre com a mensagem
    pronta), um desalinhamento manda a mensagem pra pessoa errada -- a
    planilha inteira falha em vez de sair com links errados em silêncio."""


def _validar_place_ids(nome_aba: str, linhas: list) -> None:
    faltando = [i for i, linha in enumerate(linhas, start=2) if not linha.get("place_id")]
    if faltando:
        raise PlaceIdAusenteError(
            f'Aba "{nome_aba}": {len(faltando)} linha(s) sem place_id (linhas {faltando}) -- '
            f"nenhuma planilha foi gerada."
        )


_RE_SO_DIGITOS = re.compile(r"\D")

# (coluna do link, coluna do dado que o link tem que refletir)
_PARES_LINK_DADO = (("WhatsApp", "telefone"), ("Abrir e-mail", "email"))


def _verificar_links_aba(nome_aba: str, ws, colunas) -> list:
    """Confere, linha a linha, que o hyperlink (destino do clique) de
    WhatsApp/e-mail bate com o telefone/email DA MESMA LINHA, e que o texto
    visível da célula é igual ao hyperlink -- ver `LinkInconsistenteError`
    pro defeito real que motivou esta checagem. Devolve uma lista de
    descrições de problema (vazia = nada errado)."""
    cab = list(colunas)
    if "WhatsApp" not in cab and "Abrir e-mail" not in cab:
        return []
    idx_tel = cab.index("telefone") + 1 if "telefone" in cab else None
    idx_email = cab.index("email") + 1 if "email" in cab else None
    problemas = []
    for r in range(2, ws.max_row + 1):
        for col_link, col_dado in _PARES_LINK_DADO:
            if col_link not in cab:
                continue
            idx_link = cab.index(col_link) + 1
            cel = ws.cell(r, idx_link)
            alvo = cel.hyperlink.target if cel.hyperlink else None
            if alvo is None:
                continue
            if cel.value != alvo:
                problemas.append(
                    f'Aba "{nome_aba}", linha {r}: texto de "{col_link}" difere do link '
                    f"(texto={cel.value!r}, link={alvo!r})."
                )
            if col_link == "WhatsApp" and idx_tel:
                tel = ws.cell(r, idx_tel).value
                if tel and _RE_SO_DIGITOS.sub("", str(tel)) not in alvo:
                    problemas.append(
                        f'Aba "{nome_aba}", linha {r}: link de WhatsApp não corresponde ao '
                        f"telefone da linha (telefone={tel!r}, link={alvo!r})."
                    )
                elif not tel:
                    problemas.append(
                        f'Aba "{nome_aba}", linha {r}: link de WhatsApp preenchido mas a linha '
                        f"não tem telefone."
                    )
            if col_link == "Abrir e-mail" and idx_email:
                email = ws.cell(r, idx_email).value
                if email and email not in alvo:
                    problemas.append(
                        f'Aba "{nome_aba}", linha {r}: link de e-mail não corresponde ao email '
                        f"da linha (email={email!r}, link={alvo[:60]!r})."
                    )
                elif not email:
                    problemas.append(
                        f'Aba "{nome_aba}", linha {r}: link de e-mail preenchido mas a linha não '
                        f"tem email."
                    )
    return problemas


def verificar_consistencia_links(caminho) -> dict:
    """Carrega um `.xlsx` já gravado e roda `_verificar_links_aba` em toda aba
    com coluna "WhatsApp" ou "Abrir e-mail". Devolve `{nome_aba: [problemas]}`
    (listas vazias = aba limpa). Só leitura -- não corrige nada; use pra
    conferir um arquivo já existente (comando `--verificar`, abaixo)."""
    wb = load_workbook(caminho)
    resultado = {}
    for nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
        cab = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        resultado[nome_aba] = _verificar_links_aba(nome_aba, ws, cab)
    return resultado


def diretorio_saida_padrao() -> Path:
    """Entrega INTERNA (do programa, uma planilha nova por rodada): sempre
    `comercial/_planilhas/`. `COMERCIAL_PLANILHAS_DIR` não afeta mais este
    destino -- essa env agora é só da entrega CONSOLIDADA (ver
    `diretorio_consolidada`). Override explícito continua via `--saida-dir`."""
    return DIRETORIO_SAIDA_PADRAO


def diretorio_consolidada() -> Optional[Path]:
    """Pasta da entrega CONSOLIDADA (fora do repo), lida de
    `COMERCIAL_PLANILHAS_DIR` (`comercial/.env` ou env real). `None` se não
    configurada -- a atualização da consolidada é pulada nesse caso, sem
    falhar a rodada."""
    env_loader.carregar_env()
    valor = (os.environ.get(ENV_CONSOLIDADA_DIR) or "").strip()
    return Path(valor) if valor else None


# --- Localização das fontes -----------------------------------------------


def _timestamp_do_nome(caminho: Path) -> Optional[str]:
    achado = _RE_TIMESTAMP.search(caminho.name)
    return achado.group(1) if achado else None


def localizar_csv_humano_mais_recente(diretorio: Optional[Path] = None) -> Path:
    diretorio = Path(diretorio) if diretorio else eqc.diretorio_saida_humana()
    candidatos = sorted(
        (p for p in diretorio.glob("qualificador_*.csv") if _timestamp_do_nome(p)),
        key=_timestamp_do_nome,
    )
    if not candidatos:
        raise FileNotFoundError(f"Nenhum CSV humano ('qualificador_*.csv') encontrado em {diretorio}")
    return candidatos[-1]


def ler_csv_humano(caminho: Path) -> list:
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=";"))


# --- Registro de abordados (exclusão na geração) ---------------------------


def place_ids_ja_abordados(caminho_registro: Optional[Path] = None) -> set:
    """`place_id`s já registrados em `registro_abordagens.xlsx` (delega a
    `registro_abordagens.place_ids_registrados` — uma única leitura do
    formato do registro). `caminho_registro` omitido usa o caminho padrão
    (`comercial/_abordagens/registro_abordagens.xlsx`)."""
    caminho = Path(caminho_registro) if caminho_registro else registro_abordagens.caminho_registro()
    return registro_abordagens.place_ids_registrados(caminho)


# --- Canal (decisão do diretor, 24/09/2026, quinta rodada) ------------------


def determinar_canal(telefone: Optional[str], email: Optional[str]) -> str:
    """`"whatsapp"` (celular espanhol válido, 6/7) tem prioridade; senão
    `"email"` quando há e-mail; senão `"para depois"` (nenhum canal de
    contato direto). Celular presente: só WhatsApp é usado, mesmo com
    e-mail também disponível -- "Sem celular e com e-mail => só 'Abrir
    e-mail'" implica a mesma exclusividade no sentido contrário."""
    if whatsapp_utils.normalizar_movel_espanhol(telefone):
        return "whatsapp"
    if email:
        return "email"
    return "para depois"


# --- Profissional detectado (decisão do diretor, 24/09/2026, quinta rodada) -


def profissional_detectado(nome: Optional[str]) -> str:
    """`"sim"` quando `nome` começa com Dr./Dra./Doctor/Doctora (qualquer
    caixa, com ou sem ponto), `"não"` caso contrário. NUNCA extrai o nome do
    profissional -- só a flag booleana."""
    return "sim" if nome and _RE_PROFISSIONAL.match(nome) else "não"


# --- Captura de tela (decisão do diretor, 24/09/2026, quinta rodada) --------


def link_captura(place_id: Optional[str], classe_site: Optional[str], diretorio: Path) -> tuple:
    """`(valor_da_celula, aviso)` -- link `file://` para
    `<place_id>_mobile.png` das capturas do QUALIFICADOR (`eqc.diretorio_
    capturas()`, resolvido pela raiz do LeadFarm, nunca caminho absoluto no
    código). Captura só é esperada para lead com site PRÓPRIO
    (`classe_site == "proprio"`) -- só esses passam pela fase de
    renderização do QUALIFICADOR (`main.py:fase_render`, que só mede
    `classe_site == "proprio"`). Fora disso (defeito real, 24/09/2026: o
    aviso "captura não encontrada" aparecia até em leads da Direta, que
    nunca têm captura), célula vazia e SEM aviso -- não é uma falta, é
    esperado. Para `classe_site == "proprio"`: aviso quando falta
    `place_id`, ou o arquivo não existe."""
    if classe_site != "proprio":
        return "", ""
    if not place_id:
        return "", "sem captura (place_id ausente)"
    caminho = Path(diretorio) / f"{place_id}_mobile.png"
    if caminho.is_file():
        return caminho.as_uri(), ""
    return "", f"captura não encontrada: {caminho.name}"


# --- Contexto para IA -------------------------------------------------------
#
# Texto único, montado por PROGRAMA a partir dos dados já coletados/gerados
# -- nenhuma chamada de LLM aqui. Serve para o diretor colar numa conversa
# com uma IA quando o cliente responde. Respeita LIMITE_CELULA_XLSX: só a
# seção "TEXTO DO SITE" é truncada quando o total passaria do limite (é a
# única parte de tamanho variável grande — até 4000 chars).


def _juntar_secoes(linhas_fixas_antes, texto_site_secao, linhas_fixas_depois) -> str:
    return "\n".join([*linhas_fixas_antes, f"TEXTO DO SITE: {texto_site_secao}", *linhas_fixas_depois])


def _montar_com_limite(linhas_antes, texto_site_txt, linhas_depois, limite=LIMITE_CELULA_XLSX) -> str:
    completo = _juntar_secoes(linhas_antes, texto_site_txt, linhas_depois)
    if len(completo) <= limite:
        return completo
    overhead = len(_juntar_secoes(linhas_antes, "", linhas_depois)) + len(_AVISO_TRUNCAMENTO)
    espaco = max(limite - overhead, 0)
    texto_cortado = texto_site_txt[:espaco] + _AVISO_TRUNCAMENTO
    return _juntar_secoes(linhas_antes, texto_cortado, linhas_depois)


def _linha_lead_generica(nome, nicho, cidade, telefone, site, google_maps_url) -> str:
    partes = [nome]
    if nicho:
        partes.append(nicho)
    if cidade:
        partes.append(cidade)
    partes += [telefone or "(sem telefone)", site or "(sem site)", google_maps_url or "(sem link)"]
    return "LEAD: " + " · ".join(partes)


def montar_contexto_ia_geral(linha: dict) -> str:
    """"Contexto para IA" da aba Geral: dados do lead, classe do site,
    ângulo escolhido e a mensagem sugerida."""
    linha_lead = _linha_lead_generica(
        linha.get("nome", ""), None, linha.get("cidade"), linha.get("telefone"),
        linha.get("site"), linha.get("google_maps_url"),
    )
    classe = linha.get("classe_site") or ""
    angulo = linha.get("Ângulo") or ""
    return "\n".join([
        linha_lead,
        f"CLASSE DO SITE: {classe or '(não classificado)'}",
        f"ÂNGULO: {angulo or '(nenhum)'}",
        f"MENSAGEM SUGERIDA: {linha.get('Mensagem sugerida') or '(nenhuma)'}",
    ])[:LIMITE_CELULA_XLSX]


# --- Linhas da aba Geral (pista Direta) -------------------------------------


def _linha_geral(
    lead: dict, *, apresentacao, config_validacao, assuntos_email, mensagens_angulo, diretorio_capturas,
) -> dict:
    telefone = lead.get("telefone")
    email = lead.get("email")
    numero = whatsapp_utils.normalizar_movel_espanhol(telefone)
    canal = determinar_canal(telefone, email)

    linha = dict(lead)
    linha["Canal"] = canal
    linha["Aviso"] = ""
    linha["Mensagem sugerida"] = ""
    linha["Abrir e-mail"] = ""
    linha["Profissional detectado"] = profissional_detectado(lead.get("nome"))

    angulo = angulo_mensagem.escolher_angulo_direta(lead)
    linha["Ângulo"] = angulo
    if angulo != angulo_mensagem.SEM_ANGULO:
        mensagem, entrada_derivada = angulo_mensagem.montar_mensagem_direta(
            angulo, lead, apresentacao=apresentacao, mensagens=mensagens_angulo,
        )
        resultado_validacao = angulo_mensagem.validar_mensagem_angulo(
            mensagem, angulo, entrada_derivada, nome_negocio=lead.get("nome"), config_validacao=config_validacao,
            palavras_genericas_nome=angulo_mensagem.palavras_genericas_nome(angulo_mensagem.carregar_cidade_busca()),
        )
        if not resultado_validacao.valido:
            linha["Aviso"] = "; ".join(resultado_validacao.motivos)
        else:
            linha["Mensagem sugerida"] = mensagem
            if canal == "email":
                linha["Abrir e-mail"] = email_utils.link_mailto(email, assuntos_email["direta"], mensagem)

    linha["WhatsApp"] = (
        whatsapp_utils.link_whatsapp_com_mensagem(numero, linha["Mensagem sugerida"])
        if numero and linha["Mensagem sugerida"]
        else whatsapp_utils.link_whatsapp(numero) if numero else ""
    )

    captura, aviso_captura = link_captura(lead.get("place_id"), lead.get("classe_site"), diretorio_capturas)
    linha["Captura"] = captura
    if aviso_captura:
        linha["Aviso"] = "; ".join(filter(None, [linha["Aviso"], aviso_captura]))

    linha["Contexto para IA"] = montar_contexto_ia_geral(linha)
    return linha


def montar_linhas_geral(leads_csv: list, **kwargs) -> list:
    return [_linha_geral(lead, **kwargs) for lead in leads_csv]


# --- Linhas da aba Nata (nata + candidatos_triagem) -------------------------


def _campo_evidencia_ou_vazio(lead, chave1, chave2=None):
    bloco = lead[chave1]
    campo = bloco[chave2] if chave2 else bloco
    return campo.valor_confirmado if campo.presente() else ""


def _resumo_problema_vendavel(lead) -> str:
    itens = lead["problema_vendavel"]
    return "; ".join(item["tipo"] for item in itens)


def _tem_problema_fora_do_ar(lead) -> bool:
    return any(item["tipo"] == "fora_do_ar" for item in lead["problema_vendavel"])


def _problemas_medidos_texto(lead) -> list:
    """Cada problema vendável com a medição -- a evidência CRUA do contrato
    (não a versão saneada que `entrada_estrategia` manda ao modelo): aqui é
    material de auditoria pro diretor/IA, não entrada de LLM com as mesmas
    restrições."""
    medidos = [
        f"{item['tipo']} ({item['medicao']}): {item['evidencia']}"
        for item in lead["problema_vendavel"]
    ]
    nota = lead["reputacao"]["nota_google"]
    avaliacoes = lead["reputacao"]["review_count"]
    if nota.presente() or avaliacoes.presente():
        nota_txt = nota.valor_confirmado if nota.presente() else "não confirmada"
        aval_txt = avaliacoes.valor_confirmado if avaliacoes.presente() else "não confirmadas"
        medidos.append(f"nota {nota_txt}, avaliações {aval_txt}")
    return medidos


def montar_contexto_ia_nata(lead, angulo: str, mensagem: Optional[str]) -> str:
    """"Contexto para IA" da aba Nata -- ordem fixa: LEAD, O QUE MEDIMOS,
    TEXTO DO SITE, ÂNGULO, MENSAGEM. Sem chamada de LLM: `angulo`/`mensagem`
    vêm de `angulo_mensagem.py` (MVP, 24/09/2026)."""
    identidade = lead["identidade"]
    telefone = _campo_evidencia_ou_vazio(lead, "contato", "telefone")
    site = _campo_evidencia_ou_vazio(lead, "presenca_digital", "website_url")

    linha_lead = _linha_lead_generica(
        identidade["nome"], identidade["nicho"], identidade["cidade"], telefone, site,
        identidade["google_maps_url"],
    )
    medidos = _problemas_medidos_texto(lead)
    linha_medimos = "O QUE MEDIMOS: " + ("; ".join(medidos) if medidos else "nada medido")

    texto_site_txt = entrada_estrategia.texto_do_site(lead) or "não medido"

    linha_angulo = f"ÂNGULO: {angulo}"
    linha_mensagem = f"MENSAGEM: {mensagem or '(nenhuma -- sem ângulo)'}"

    return _montar_com_limite([linha_lead, linha_medimos], texto_site_txt, [linha_angulo, linha_mensagem])


def _linha_nata(
    lead, pista: str, leads_do_lote: list, *,
    apresentacao, config_validacao, assuntos_email, regras_angulo, mensagens_angulo, diretorio_capturas,
) -> dict:
    telefone = _campo_evidencia_ou_vazio(lead, "contato", "telefone")
    email = _campo_evidencia_ou_vazio(lead, "contato", "email")
    numero = whatsapp_utils.normalizar_movel_espanhol(telefone)
    canal = determinar_canal(telefone, email)

    linha = {
        "pista": pista,
        "place_id": lead["place_id"],
        "nome": lead["identidade"]["nome"],
        "telefone": telefone,
        "WhatsApp": "",
        "email": email,
        "Abrir e-mail": "",
        "Canal": canal,
        "problema vendável": _resumo_problema_vendavel(lead),
        "Ângulo": "",
        "mensagem": "",
        "Status": "",
        "conferir_antes_de_enviar": "sim" if _tem_problema_fora_do_ar(lead) else "não",
        "Profissional detectado": profissional_detectado(lead["identidade"]["nome"]),
        "Captura": "",
        "site": _campo_evidencia_ou_vazio(lead, "presenca_digital", "website_url"),
        "google_maps_url": lead["identidade"]["google_maps_url"] or "",
        "texto_site": entrada_estrategia.texto_do_site(lead) or "",
        "Aviso": "",
    }
    for c in COLUNAS_DIRETOR:
        linha[c] = ""

    angulo = angulo_mensagem.escolher_angulo_nata(lead, leads_do_lote, regras_angulo)
    linha["Ângulo"] = angulo
    mensagem_final = None
    if angulo == angulo_mensagem.SEM_ANGULO:
        linha["Status"] = STATUS_SEM_MENSAGEM
    else:
        mensagem, entrada_derivada = angulo_mensagem.montar_mensagem_nata(
            angulo, lead, leads_do_lote, apresentacao=apresentacao, regras=regras_angulo,
            mensagens=mensagens_angulo,
        )
        resultado_validacao = angulo_mensagem.validar_mensagem_angulo(
            mensagem, angulo, entrada_derivada, nome_negocio=lead["identidade"]["nome"],
            config_validacao=config_validacao,
            palavras_genericas_nome=angulo_mensagem.palavras_genericas_nome(
                lead["identidade"]["nicho"], angulo_mensagem.carregar_cidade_busca(),
            ),
        )
        if not resultado_validacao.valido:
            linha["Aviso"] = "; ".join(resultado_validacao.motivos)
            linha["Status"] = STATUS_SEM_MENSAGEM
        else:
            linha["mensagem"] = mensagem
            mensagem_final = mensagem
            if canal == "email":
                linha["Abrir e-mail"] = email_utils.link_mailto(email, assuntos_email["nata"], mensagem)

    linha["WhatsApp"] = (
        whatsapp_utils.link_whatsapp_com_mensagem(numero, linha["mensagem"])
        if numero and linha["mensagem"]
        else whatsapp_utils.link_whatsapp(numero) if numero else ""
    )

    captura, aviso_captura = link_captura(lead["place_id"], lead["classe_site"], diretorio_capturas)
    linha["Captura"] = captura
    if aviso_captura:
        linha["Aviso"] = "; ".join(filter(None, [linha["Aviso"], aviso_captura]))

    linha["Contexto para IA"] = montar_contexto_ia_nata(lead, angulo, mensagem_final)
    return linha


def montar_linhas_nata(nata: list, candidatos_triagem: list, descartados: list, **kwargs) -> list:
    """`nata` e `candidatos_triagem` viram linhas (coluna `pista` distingue
    as duas); `descartados` entra só como pool de concorrentes do ângulo
    `reputacao` (`angulo_mensagem.escolher_angulo_nata`), nunca vira linha
    própria."""
    leads_do_lote = list(nata) + list(candidatos_triagem) + list(descartados)
    linhas = [_linha_nata(lead, "nata", leads_do_lote, **kwargs) for lead in nata]
    linhas += [
        _linha_nata(lead, "candidatos_triagem", leads_do_lote, **kwargs)
        for lead in candidatos_triagem
    ]
    return linhas


# --- Aba "Como usar" ----------------------------------------------------------


def _escrever_aba_como_usar(wb) -> None:
    ws = wb.create_sheet("Como usar")
    ws.append(["Regra", "O que fazer"])
    for cel in ws[1]:
        cel.font = Font(bold=True)
    for titulo, texto in _COMO_USAR_TEXTO:
        ws.append([titulo, texto])
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100


# --- Escrita da planilha -----------------------------------------------------


def _sem_caracteres_ilegais(valor):
    """Remove caracteres de controle que o Excel recusa
    (`openpyxl.cell.cell.ILLEGAL_CHARACTERS_RE` -- ex.: `\\x0b`, `\\x1f`) de
    um valor de célula. Defeito real, 25/09/2026: o `texto_site` de um lead
    real trazia um caractere de controle e `IllegalCharacterError` impedia a
    planilha inteira de ser gerada -- nenhuma linha chegava a ser gravada.
    Ponto único de limpeza: toda célula de texto passa por aqui antes de
    `ws.append`, nunca em cada lugar que monta uma linha."""
    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)
    return valor


def _acrescentar_linhas(ws, colunas, linhas, *, cabecalho: bool):
    """Só grava: cabeçalho (quando `cabecalho=True`) + uma linha por item de
    `linhas`. Sem estilo -- isso é `_restilizar_aba`, separado pra poder
    reaplicar o estilo depois de um `ws.append` feito em outra rodada (caso
    da planilha consolidada, que é reaberta e acrescentada, nunca recriada).

    `cabecalho` é decidido pelo CHAMADOR (aba nova ou não) em vez de sondado
    aqui: ler `ws.cell(1, 1).value` numa aba vazia já materializa a célula no
    openpyxl e faz o próximo `ws.append` pular pra linha 2 -- defeito real,
    achado nesta tarefa (01/10/2026), que deixava a aba com uma linha 1 em
    branco e o cabeçalho na linha 2."""
    if cabecalho:
        ws.append([_sem_caracteres_ilegais(c) for c in colunas])
        for cel in ws[1]:
            cel.font = Font(bold=True)
        ws.freeze_panes = "A2"

    for linha in linhas:
        ws.append([_sem_caracteres_ilegais(linha.get(c, "")) for c in colunas])


def _restilizar_aba(ws, colunas, *, colunas_link=(), colunas_wrap=(), opcoes_funil=None):
    """(Re)aplica filtro, validação de funil, links, quebra de linha, largura
    de coluna e a coluna `place_id` oculta sobre o estado ATUAL da aba (todas
    as linhas, não só as que acabaram de entrar) -- pode ser chamado tanto
    numa aba recém-criada quanto numa reaberta e acrescentada."""
    ultima_linha = ws.max_row
    if ultima_linha >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{max(ultima_linha, 1)}"

    if opcoes_funil:
        registro_abordagens.aplicar_validacoes_funil(ws, colunas, opcoes_funil, ultima_linha)

    for nome_col in colunas_link:
        idx = colunas.index(nome_col) + 1
        for r in range(2, ultima_linha + 1):
            cel = ws.cell(row=r, column=idx)
            if cel.value:
                cel.hyperlink = cel.value
                cel.font = Font(color="0563C1", underline="single")

    for nome_col in colunas_wrap:
        idx = colunas.index(nome_col) + 1
        for r in range(2, ultima_linha + 1):
            ws.cell(row=r, column=idx).alignment = Alignment(wrap_text=True, vertical="top")

    for i, nome_col in enumerate(colunas, start=1):
        maior = len(nome_col)
        for r in range(2, ultima_linha + 1):
            maior = max(maior, len(str(ws.cell(r, i).value or "")))
        ws.column_dimensions[get_column_letter(i)].width = min(max(maior + 2, 10), 60)

    if "place_id" in colunas:
        ws.column_dimensions[get_column_letter(colunas.index("place_id") + 1)].hidden = True


def _escrever_aba(ws, colunas, linhas, *, colunas_link=(), colunas_wrap=(), opcoes_funil=None):
    _acrescentar_linhas(ws, colunas, linhas, cabecalho=True)
    _restilizar_aba(ws, colunas, colunas_link=colunas_link, colunas_wrap=colunas_wrap, opcoes_funil=opcoes_funil)


def _place_ids_na_aba(ws, colunas) -> set:
    """`place_id`s já presentes numa aba (linhas de dados, sem o cabeçalho).
    Usado pra não duplicar lead na planilha consolidada entre rodadas."""
    if "place_id" not in colunas or ws.max_row < 2:
        return set()
    idx = colunas.index("place_id") + 1
    return {
        str(ws.cell(r, idx).value) for r in range(2, ws.max_row + 1) if ws.cell(r, idx).value
    }


def atualizar_planilha_consolidada(
    *,
    linhas_geral: list,
    linhas_nata: list,
    opcoes_funil: Optional[dict] = None,
    diretorio: Optional[Path] = None,
) -> dict:
    """Mantém a entrega CONSOLIDADA do diretor: um único `.xlsx` fora do
    repo, com só as abas Geral e Nata (sem "Como usar"), que cada rodada
    ABRE e ACRESCENTA -- nunca recria nem sobrescreve uma rodada anterior
    (decisão do diretor, 01/10/2026). Dedup por `place_id`: lead já presente
    na aba não entra de novo.

    `diretorio` omitido = lê `COMERCIAL_PLANILHAS_DIR`. Pasta não configurada
    = cópia pulada, a rodada NUNCA falha por causa disto (mesmo espírito de
    `saida_humana.copiar_conveniencia` do QUALIFICADOR) -- devolve
    `{"pulado": True, "motivo": ...}`."""
    destino_dir = diretorio if diretorio is not None else diretorio_consolidada()
    if destino_dir is None:
        return {"pulado": True, "motivo": f"{ENV_CONSOLIDADA_DIR} não definida -- consolidada pulada."}

    caminho = Path(destino_dir) / NOME_PLANILHA_CONSOLIDADA
    try:
        destino_dir = Path(destino_dir)
        destino_dir.mkdir(parents=True, exist_ok=True)

        if caminho.exists():
            wb = load_workbook(caminho)
        else:
            # Workbook() já vem com uma aba padrão ("Sheet"); removê-la ANTES
            # de checar sheetnames -- renomear essa aba padrão pra "Geral"
            # (em vez de removê-la) faria "Geral" já constar em sheetnames
            # antes da checagem abaixo, e o cabeçalho seria pulado também
            # numa planilha nova (defeito real, achado nesta tarefa).
            wb = Workbook()
            wb.remove(wb.active)

        geral_nova = "Geral" not in wb.sheetnames
        ws_geral = wb.create_sheet("Geral") if geral_nova else wb["Geral"]
        nata_nova = "Nata" not in wb.sheetnames
        ws_nata = wb.create_sheet("Nata") if nata_nova else wb["Nata"]

        ja_geral = _place_ids_na_aba(ws_geral, COLUNAS_GERAL)
        ja_nata = _place_ids_na_aba(ws_nata, COLUNAS_NATA)
        novas_geral = [l for l in linhas_geral if str(l.get("place_id")) not in ja_geral]
        novas_nata = [l for l in linhas_nata if str(l.get("place_id")) not in ja_nata]

        _acrescentar_linhas(ws_geral, COLUNAS_GERAL, novas_geral, cabecalho=geral_nova)
        _acrescentar_linhas(ws_nata, COLUNAS_NATA, novas_nata, cabecalho=nata_nova)
        _restilizar_aba(
            ws_geral, COLUNAS_GERAL,
            colunas_link=("WhatsApp", "Abrir e-mail", "Captura"),
            colunas_wrap=("Mensagem sugerida", "Contexto para IA"),
            opcoes_funil=opcoes_funil,
        )
        _restilizar_aba(
            ws_nata, COLUNAS_NATA,
            colunas_link=("WhatsApp", "Abrir e-mail", "Captura"),
            colunas_wrap=("mensagem", "texto_site", "Contexto para IA"),
            opcoes_funil=opcoes_funil,
        )

        # Checagem antes de salvar -- a consolidada é reaberta e acrescentada
        # rodada após rodada, então um desalinhamento de hyperlink (ex.:
        # `delete_rows` numa limpeza manual) que passasse batido aqui
        # contaminaria toda rodada futura. Não captura `LinkInconsistenteError`:
        # propaga pra fora, SEM salvar -- o arquivo em disco fica como estava.
        problemas = _verificar_links_aba("Geral", ws_geral, COLUNAS_GERAL) + _verificar_links_aba(
            "Nata", ws_nata, COLUNAS_NATA
        )
        if problemas:
            raise LinkInconsistenteError(
                f"{len(problemas)} link(s) desalinhado(s) na planilha consolidada -- nada foi "
                f"salvo, o arquivo em {caminho} continua como estava:\n" + "\n".join(problemas)
            )

        wb.save(caminho)
    except OSError as e:
        return {"pulado": True, "motivo": f"consolidada falhou ({e}) -- a rodada segue normalmente."}

    return {
        "pulado": False,
        "caminho": caminho,
        "novos_geral": len(novas_geral),
        "novos_nata": len(novas_nata),
    }


def _caminho_planilha(saida_dir: Path, agora: datetime) -> Path:
    return Path(saida_dir) / f"planilha_envio_{agora.strftime('%Y%m%d-%H%M%S')}.xlsx"


def gerar_planilha(
    *,
    caminho_csv_humano: Optional[Path] = None,
    caminho_lote: Optional[Path] = None,
    caminho_registro: Optional[Path] = None,
    saida_dir: Optional[Path] = None,
    agora: Optional[datetime] = None,
    apresentacao: Optional[str] = None,
    config_validacao: Optional[dict] = None,
    assuntos_email: Optional[dict] = None,
    opcoes_funil: Optional[dict] = None,
    regras_angulo: Optional[dict] = None,
    mensagens_angulo: Optional[dict] = None,
    diretorio_capturas: Optional[Path] = None,
    consolidada_dir: Optional[Path] = None,
) -> dict:
    """Gera `planilha_envio_<timestamp>.xlsx` (entrega interna, abas Geral,
    Nata e "Como usar") e atualiza a entrega CONSOLIDADA (fora do repo, só
    Geral e Nata, acrescentada -- nunca recriada). Nunca sobrescreve a
    interna (`PlanilhaJaExisteError`). Exclui, das duas, quem já está no
    registro de abordados. Devolve
    `{"caminho", "excluidos_geral", "excluidos_nata", "consolidada"}`.

    `consolidada_dir` omitido = lê `COMERCIAL_PLANILHAS_DIR` (ver
    `diretorio_consolidada`); passe um diretório explícito em teste pra não
    tocar a pasta real do diretor."""
    saida_dir = Path(saida_dir) if saida_dir else diretorio_saida_padrao()
    agora = agora or datetime.now(timezone.utc)
    caminho = _caminho_planilha(saida_dir, agora)
    if caminho.exists():
        raise PlanilhaJaExisteError(f"{caminho} já existe — nunca sobrescrever uma planilha anterior.")

    apresentacao = apresentacao if apresentacao is not None else angulo_mensagem.carregar_apresentacao()
    config_validacao = config_validacao or carregar_config_validacao()
    assuntos_email = assuntos_email if assuntos_email is not None else email_utils.carregar_assuntos_email()
    opcoes_funil = opcoes_funil if opcoes_funil is not None else registro_abordagens.carregar_opcoes_funil()
    regras_angulo = regras_angulo if regras_angulo is not None else angulo_mensagem.carregar_regras_angulo()
    mensagens_angulo = mensagens_angulo if mensagens_angulo is not None else angulo_mensagem.carregar_mensagens_angulo()
    diretorio_capturas = Path(diretorio_capturas) if diretorio_capturas else eqc.diretorio_capturas()

    caminho_csv = Path(caminho_csv_humano) if caminho_csv_humano else localizar_csv_humano_mais_recente()
    leads_csv = ler_csv_humano(caminho_csv)

    lote = contrato_loader.carregar_lote(caminho=caminho_lote)

    ja_abordados = place_ids_ja_abordados(caminho_registro)
    excluidos_geral = [l for l in leads_csv if l.get("place_id") in ja_abordados]
    leads_csv = [l for l in leads_csv if l.get("place_id") not in ja_abordados]
    nata_excluida = [l for l in lote.nata if l["place_id"] in ja_abordados]
    nata_incluida = [l for l in lote.nata if l["place_id"] not in ja_abordados]
    candidatos_excluidos = [l for l in lote.candidatos_triagem if l["place_id"] in ja_abordados]
    candidatos_incluidos = [l for l in lote.candidatos_triagem if l["place_id"] not in ja_abordados]

    kwargs_comuns = dict(
        apresentacao=apresentacao, config_validacao=config_validacao, assuntos_email=assuntos_email,
        mensagens_angulo=mensagens_angulo, diretorio_capturas=diretorio_capturas,
    )

    linhas_geral = montar_linhas_geral(leads_csv, **kwargs_comuns)
    linhas_nata = montar_linhas_nata(
        nata_incluida, candidatos_incluidos, lote.descartados, regras_angulo=regras_angulo,
        **kwargs_comuns,
    )

    _validar_place_ids("Geral", linhas_geral)
    _validar_place_ids("Nata", linhas_nata)

    wb = Workbook()
    ws_geral = wb.active
    ws_geral.title = "Geral"
    _escrever_aba(
        ws_geral, COLUNAS_GERAL, linhas_geral,
        colunas_link=("WhatsApp", "Abrir e-mail", "Captura"),
        colunas_wrap=("Mensagem sugerida", "Contexto para IA"),
        opcoes_funil=opcoes_funil,
    )
    ws_nata = wb.create_sheet("Nata")
    _escrever_aba(
        ws_nata, COLUNAS_NATA, linhas_nata,
        colunas_link=("WhatsApp", "Abrir e-mail", "Captura"),
        colunas_wrap=("mensagem", "texto_site", "Contexto para IA"),
        opcoes_funil=opcoes_funil,
    )
    _escrever_aba_como_usar(wb)

    problemas = _verificar_links_aba("Geral", ws_geral, COLUNAS_GERAL) + _verificar_links_aba(
        "Nata", ws_nata, COLUNAS_NATA
    )
    if problemas:
        raise LinkInconsistenteError(
            f"{len(problemas)} link(s) desalinhado(s) -- nenhuma planilha foi gerada:\n"
            + "\n".join(problemas)
        )

    saida_dir.mkdir(parents=True, exist_ok=True)
    wb.save(caminho)

    consolidada = atualizar_planilha_consolidada(
        linhas_geral=linhas_geral, linhas_nata=linhas_nata, opcoes_funil=opcoes_funil,
        diretorio=consolidada_dir,
    )

    return {
        "caminho": caminho,
        "excluidos_geral": len(excluidos_geral),
        "excluidos_nata": len(nata_excluida) + len(candidatos_excluidos),
        "consolidada": consolidada,
    }


def main(argv: Optional[list] = None) -> int:
    ap = argparse.ArgumentParser(
        description="Etapa E2b — planilha de envio em .xlsx (abas Geral, Nata e Como usar)."
    )
    ap.add_argument(
        "--saida-dir", type=Path, default=None,
        help="Diretório de saída interna (default: comercial/_planilhas/).",
    )
    ap.add_argument(
        "--verificar", type=Path, default=None, metavar="ARQUIVO",
        help="Só confere os hyperlinks de um .xlsx já existente (WhatsApp/e-mail batendo com "
        "telefone/email da própria linha); não gera planilha nova. Exit 0 = tudo certo, 1 = "
        "achou desalinhamento.",
    )
    args = ap.parse_args(argv)

    if args.verificar:
        resultado = verificar_consistencia_links(args.verificar)
        total = sum(len(v) for v in resultado.values())
        for nome_aba, problemas in resultado.items():
            print(f'Aba "{nome_aba}": {len(problemas)} problema(s).')
            for p in problemas:
                print(f"  {p}")
        print("OK, nenhum link desalinhado." if total == 0 else f"TOTAL: {total} problema(s).")
        return 0 if total == 0 else 1

    resultado = gerar_planilha(saida_dir=args.saida_dir)

    print(
        f"OK: planilha interna gerada em {resultado['caminho']} "
        f"({resultado['excluidos_geral']} excluído(s) na Geral, "
        f"{resultado['excluidos_nata']} excluído(s) na Nata -- já abordados)."
    )
    consolidada = resultado["consolidada"]
    if consolidada.get("pulado"):
        print(f"[Consolidada] {consolidada['motivo']}")
    else:
        print(
            f"[Consolidada] {consolidada['caminho']}: "
            f"+{consolidada['novos_geral']} novo(s) na Geral, +{consolidada['novos_nata']} novo(s) na Nata."
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
