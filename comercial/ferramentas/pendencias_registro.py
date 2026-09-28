"""pendencias_registro.py — o que a planilha de envio tem e o registro ainda não.

Só LEITURA: nunca escreve na planilha nem no registro. Compara uma planilha
de envio preenchida pelo diretor com `registro_abordagens.xlsx` e lista:

1. **envio não registrado** — linha com "Data 1º contato" preenchida ou
   "Decidi não enviar" = "sim" cujo `place_id` ainda não está no registro
   (resolve com `registro_abordagens.py <planilha>`);
2. **acompanhamento divergente** — `place_id` já registrado cujo Etapa,
   Resultado, "Enviada como" ou "Decidi não enviar" na planilha difere do
   registro. O registro só ACRESCENTA (decisão do diretor, 24/09/2026): essa
   diferença não é gravada por nenhum comando -- é listada para o diretor
   decidir.

Decisão do diretor, 28/09/2026: a conferência de resultados vira hábito por
MECANISMO (este comando, com exit code), não por memória de sessão. Exit 0 =
nada pendente; 1 = há pendência; 2 = arquivo ausente ou ilegível.
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

try:
    import registro_abordagens
except ModuleNotFoundError:  # pragma: no cover - bootstrap de sys.path
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import registro_abordagens

# Campo do dict de `registro_abordagens._linhas_enviadas` -> coluna do registro.
CAMPOS_ACOMPANHAMENTO = (
    ("etapa", "Etapa"),
    ("resultado", "Resultado"),
    ("mensagem_status", "Enviada como"),
    ("decidiu_nao_enviar", "Decidi não enviar"),
)

ENV_PLANILHAS_DIR = "COMERCIAL_PLANILHAS_DIR"


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def ler_registro(caminho: Path) -> dict:
    """`{place_id: {coluna: valor}}` do registro. Ausente = `{}`."""
    caminho = Path(caminho)
    if not caminho.is_file():
        return {}
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not linhas or "place_id" not in linhas[0]:
        return {}
    cabecalho = linhas[0]
    registro = {}
    for linha in linhas[1:]:
        dados = dict(zip(cabecalho, linha))
        if dados.get("place_id"):
            registro[dados["place_id"]] = dados
    return registro


def pendencias(caminho_planilha: Path, caminho_registro: Path) -> dict:
    """`{"nao_registrados": [...], "divergentes": [...]}`. Cada divergência
    traz `place_id`, `nome`, `coluna`, `no_registro`, `na_planilha`."""
    linhas = registro_abordagens._linhas_enviadas(Path(caminho_planilha))
    registro = ler_registro(caminho_registro)
    nao_registrados = []
    divergentes = []
    for linha in linhas:
        pid = linha["place_id"]
        if pid not in registro:
            nao_registrados.append({"place_id": pid, "nome": linha.get("nome", "")})
            continue
        no_registro = registro[pid]
        for campo, coluna in CAMPOS_ACOMPANHAMENTO:
            planilha = _texto(linha.get(campo))
            gravado = _texto(no_registro.get(coluna))
            if planilha != gravado:
                divergentes.append({
                    "place_id": pid, "nome": linha.get("nome", ""), "coluna": coluna,
                    "no_registro": gravado, "na_planilha": planilha,
                })
    return {"nao_registrados": nao_registrados, "divergentes": divergentes}


def planilha_mais_recente(diretorio: Path) -> Optional[Path]:
    candidatos = sorted(Path(diretorio).glob("planilha_envio_*.xlsx"))
    return candidatos[-1] if candidatos else None


def main(argv: Optional[list] = None) -> int:
    ap = argparse.ArgumentParser(description="Lista o que a planilha de envio tem e o registro ainda não (só leitura).")
    ap.add_argument("planilha", nargs="?", type=Path,
                    help=f"Planilha preenchida. Omitida: a mais recente em ${ENV_PLANILHAS_DIR} ou comercial/_planilhas/.")
    ap.add_argument("--registro", type=Path, default=None, help="Caminho do registro (default: comercial/_abordagens/).")
    args = ap.parse_args(argv)

    planilha = args.planilha
    if planilha is None:
        diretorio = Path(os.environ.get(ENV_PLANILHAS_DIR) or Path(__file__).resolve().parent.parent / "_planilhas")
        planilha = planilha_mais_recente(diretorio)
        if planilha is None:
            print(f"ERRO: nenhuma planilha_envio_*.xlsx em {diretorio}")
            return 2
    if not Path(planilha).is_file():
        print(f"ERRO: planilha não encontrada: {planilha}")
        return 2
    caminho_reg = args.registro or registro_abordagens.caminho_registro()

    try:
        r = pendencias(planilha, caminho_reg)
    except (registro_abordagens.RegistroSemPlaceIdError, OSError) as e:
        print(f"ERRO: {e}")
        return 2

    print(f"Planilha: {planilha}")
    print(f"Registro: {caminho_reg}")
    if not r["nao_registrados"] and not r["divergentes"]:
        print("OK: nada pendente -- o registro está em dia com a planilha.")
        return 0
    if r["nao_registrados"]:
        print(f"\n{len(r['nao_registrados'])} envio(s) NÃO registrado(s) -- rode registro_abordagens.py com esta planilha:")
        for p in r["nao_registrados"]:
            print(f"  - {p['nome']} ({p['place_id']})")
    if r["divergentes"]:
        print(f"\n{len(r['divergentes'])} diferença(s) de acompanhamento em leads já registrados "
              f"(o registro só acrescenta -- decidir com o diretor):")
        for d in r["divergentes"]:
            print(f"  - {d['nome']}: {d['coluna']} registro={d['no_registro'] or '(vazio)'} "
                  f"planilha={d['na_planilha'] or '(vazio)'}")
    return 1


if __name__ == "__main__":
    sys.exit(main())
