"""Coluna "Atendimento WhatsApp" (decisão do diretor, 12/10/2026): última
coluna das abas Geral e Nata, com lista de escolha de só cinco valores;
gravada no registro de abordagens pelo mesmo mecanismo do funil por eventos
(preencher ou mudar = linha nova com `evento` = "atendimento"; o mesmo valor
de novo não grava nada; o registro só acrescenta); consulta só de leitura
`--atendimento`; definição dos valores na aba "Como usar". Dados fictícios,
tudo em tmp_path; cada regra com o caso que grava e o controle que não grava."""

import datetime
import hashlib

import pytest
from openpyxl import Workbook, load_workbook

import pendencias_registro as pr
import planilha_envio as pe
import registro_abordagens as ra

_VALORES = ("humano", "auto-resposta simples", "menu/bot", "agente IA", "sem resposta (48h)")
_FUNIL = {"etapa": ["1º contato"], "resultado": ["sem resposta", "resposta positiva", "pediu exemplo"],
          "mensagem": ["enviada como está"], "decidi_nao_enviar": ["sim"], "atendimento_whatsapp": list(_VALORES)}
_CAB = ["place_id", "nome", "telefone", "Data 1º contato", "Etapa", "Resultado", "Decidi não enviar", "Canal",
        "Ângulo", "campanha", "nicho", "pista", "cidade_conferida", "Atendimento WhatsApp"]


def _planilha(caminho, linhas, abas=("Geral",)):
    wb = Workbook()
    wb.remove(wb.active)
    for aba in abas:
        ws = wb.create_sheet(aba)
        ws.append(_CAB)
        for l in linhas:
            ws.append([l.get(c) for c in _CAB])
    wb.save(caminho)
    return caminho


def _lead(pid="p1", **over):
    base = {"place_id": pid, "nome": f"Estudio Ficticio {pid}", "telefone": "600000101",
            "Data 1º contato": "2026-10-09", "Etapa": "1º contato", "Canal": "whatsapp", "Ângulo": "sem_site",
            "campanha": "arquitectos-2026-10", "nicho": "Arquitectos", "pista": "direta",
            "cidade_conferida": "Murcia"}
    base.update(over)
    return base


def _at(valor, pid="p1", **over):
    return _lead(pid, **{"Atendimento WhatsApp": valor}, **over)


def _registrar(tmp_path, planilha, dia=12):
    agora = datetime.datetime(2026, 10, dia, 12, 0, tzinfo=datetime.timezone.utc)
    return ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_FUNIL,
                               agora=agora)


def _linhas(tmp_path):
    wb = load_workbook(ra.caminho_registro(tmp_path / "_abordagens"), read_only=True)
    try:
        linhas = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    return [dict(zip(linhas[0], l)) for l in linhas[1:]]


def _hash(caminho):
    return hashlib.sha256(caminho.read_bytes()).hexdigest()


# --- planilha: coluna no fim, lista de escolha, Como usar ------------------------------


def test_config_real_tem_exatamente_os_cinco_valores_e_e_obrigatoria(tmp_path):
    assert tuple(ra.carregar_opcoes_funil()["atendimento_whatsapp"]) == _VALORES
    import json
    sem = {k: v for k, v in json.loads(ra.CAMINHO_FUNIL_PADRAO.read_text(encoding="utf-8")).items()
           if k != "atendimento_whatsapp"}
    caminho = tmp_path / "funil.json"
    caminho.write_text(json.dumps(sem), encoding="utf-8")
    with pytest.raises(ra.ConfigFunilInvalidaError, match="atendimento_whatsapp"):
        ra.carregar_opcoes_funil(caminho)


@pytest.mark.parametrize("colunas", [pe.COLUNAS_GERAL, pe.COLUNAS_NATA], ids=["Geral", "Nata"])
def test_coluna_no_fim_com_menu_de_so_cinco_valores(colunas):
    assert colunas[-1] == "Atendimento WhatsApp" and colunas.count("Atendimento WhatsApp") == 1
    wb = Workbook()
    ws = wb.active
    pe._escrever_aba(ws, colunas, [{"place_id": "p1"}, {"place_id": "p2"}],
                     opcoes_funil=ra.carregar_opcoes_funil())
    letra = ws.cell(1, len(colunas)).column_letter
    validacoes = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref) == f"{letra}2:{letra}3"]
    assert len(validacoes) == 1
    assert validacoes[0].type == "list"
    assert validacoes[0].formula1 == '"' + ",".join(_VALORES) + '"'


def test_como_usar_define_cada_valor_e_a_regra_do_primeiro():
    textos = dict(pe._COMO_USAR_TEXTO)
    texto = textos["Atendimento WhatsApp"]
    for valor in _VALORES:
        assert f'"{valor}"' in texto, valor
    assert "Marque o que chegar PRIMEIRO" in texto
    assert 'auto-resposta seguida de uma pessoa = "auto-resposta simples"' in texto


# --- registro: mesmo mecanismo do funil por eventos -------------------------------------


def test_preencher_e_mudar_viram_eventos_e_o_mesmo_valor_nao_duplica(tmp_path):
    assert _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)["acrescentados"] == 1
    r = _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_at("auto-resposta simples")]), 10)
    assert (r["eventos_atendimento"], r["eventos_resultado"], r["ja_existentes"]) == (1, 0, 0)
    # o mesmo valor de novo: nada gravado (controle)
    r = _registrar(tmp_path, _planilha(tmp_path / "c.xlsx", [_at("auto-resposta simples")]), 11)
    assert (r["eventos"], r["ja_existentes"], r["backup"]) == (0, 1, None)
    r = _registrar(tmp_path, _planilha(tmp_path / "d.xlsx", [_at("humano")]), 12)
    assert r["eventos_atendimento"] == 1
    assert [(l["evento"], l["Atendimento WhatsApp"], l["registrado_em"][:10]) for l in _linhas(tmp_path)] == [
        ("envio", None, "2026-10-09"),
        ("atendimento", "auto-resposta simples", "2026-10-10"),
        ("atendimento", "humano", "2026-10-12"),
    ]


def test_registro_so_acrescenta_linhas_antigas_ficam_iguais(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    antes = _linhas(tmp_path)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_at("menu/bot")]), 10)
    depois = _linhas(tmp_path)
    assert depois[: len(antes)] == antes
    assert len(depois) == len(antes) + 1


def test_valor_apagado_nao_vira_evento_nem_apaga_o_gravado(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_at("agente IA")]), 9)
    r = _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_at(None)]), 10)
    assert r["eventos"] == 0
    # um Resultado novo com o campo vazio na planilha não "zera" o atendimento gravado
    _registrar(tmp_path, _planilha(tmp_path / "c.xlsx", [_at(None, Resultado="resposta positiva")]), 11)
    assert ra.ultimo_atendimento(_linhas(tmp_path)) == "agente IA"
    r = _registrar(tmp_path, _planilha(tmp_path / "d.xlsx", [_at("agente IA", Resultado="resposta positiva")]), 12)
    assert r["eventos"] == 0, "o mesmo valor de antes não é mudança"


def test_lead_novo_com_o_campo_preenchido_entra_numa_linha_so(tmp_path):
    r = _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_at("humano", Resultado="resposta positiva")]), 9)
    assert (r["acrescentados"], r["eventos"]) == (1, 0)
    assert [(l["evento"], l["Atendimento WhatsApp"]) for l in _linhas(tmp_path)] == [("envio", "humano")]


def test_resultado_e_atendimento_mudando_juntos_sao_duas_linhas(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    r = _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_at("menu/bot", Resultado="pediu exemplo")]), 10)
    assert (r["eventos_resultado"], r["eventos_atendimento"]) == (1, 1)
    assert [(l["evento"], l["Resultado"], l["Atendimento WhatsApp"]) for l in _linhas(tmp_path)] == [
        ("envio", None, None),
        ("resultado", "pediu exemplo", None),
        ("atendimento", "pediu exemplo", "menu/bot"),
    ]


def test_o_mesmo_lead_em_duas_abas_gera_um_evento_so(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    r = _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_at("humano")], abas=("Geral", "Nata")), 10)
    assert (r["eventos_atendimento"], r["ja_existentes"]) == (1, 1)
    assert len(_linhas(tmp_path)) == 2


def test_funil_mostra_o_passo_de_atendimento(tmp_path, capsys):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_at("sem resposta (48h)")]), 11)
    funil = ra.funil_por_lead(ra.caminho_registro(tmp_path / "_abordagens"))
    assert [p["evento"] for p in funil["p1"]["passos"]] == ["envio", "atendimento"]
    ra.imprimir_funil(funil, tuple(_FUNIL["resultado"]), "p1")
    assert "2026-10-11 atendimento: sem resposta (48h)" in capsys.readouterr().out


# --- pendências: atendimento não registrado é evento; apagado é divergência ------------


def test_pendencias_lista_atendimento_nao_registrado_e_apagado(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_at("humano"), _lead("p2")]), 9)
    registro = ra.caminho_registro(tmp_path / "_abordagens")
    planilha = _planilha(tmp_path / "b.xlsx", [_at(None), _at("menu/bot", "p2")])
    r = pr.pendencias(planilha, registro)
    assert [(e["place_id"], e["coluna"], e["na_planilha"]) for e in r["eventos_nao_registrados"]] == [
        ("p2", "Atendimento WhatsApp", "menu/bot")]
    assert [(d["place_id"], d["coluna"], d["no_registro"]) for d in r["divergentes"]] == [
        ("p1", "Atendimento WhatsApp", "humano")]
    _registrar(tmp_path, planilha, 10)
    r = pr.pendencias(planilha, registro)
    assert r["eventos_nao_registrados"] == []  # resolvido pelo registro


# --- consulta --atendimento (só leitura) -----------------------------------------------


def _registro_de_exemplo(tmp_path):
    linhas = [
        _at("humano", "p1"),
        _at("menu/bot", "p2", nicho="Psicólogos", pista="nata", cidade_conferida="València"),
        _at("humano", "p3", nicho="Psicólogos", pista="nata", cidade_conferida="València"),
        _lead("p4", nicho="Abogados"),                                   # enviado, sem o campo
        _lead("p5", **{"Data 1º contato": None, "Decidi não enviar": "sim"}),   # nunca enviado
    ]
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", linhas), 9)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_at("agente IA", "p1")]), 10)  # mudou depois
    return ra.caminho_registro(tmp_path / "_abordagens")


def test_resumo_conta_por_tipo_e_cruza_com_nicho_pista_e_cidade(tmp_path):
    r = ra.resumo_atendimento(_registro_de_exemplo(tmp_path), _VALORES)
    assert r["total"] == 4, "p5 (Decidi não enviar, sem envio) fica fora"
    assert r["por_tipo"] == {"humano": 1, "auto-resposta simples": 0, "menu/bot": 1, "agente IA": 1,
                             "sem resposta (48h)": 0, ra.SEM_ATENDIMENTO: 1}
    assert r["por_nicho"] == {"Arquitectos": {"agente IA": 1}, "Psicólogos": {"menu/bot": 1, "humano": 1},
                              "Abogados": {ra.SEM_ATENDIMENTO: 1}}
    assert r["por_pista"] == {"direta": {"agente IA": 1, ra.SEM_ATENDIMENTO: 1},
                              "nata": {"menu/bot": 1, "humano": 1}}
    assert r["por_cidade"] == {"Murcia": {"agente IA": 1, ra.SEM_ATENDIMENTO: 1},
                               "València": {"menu/bot": 1, "humano": 1}}
    assert r["sem_campo"] == [("Estudio Ficticio p4", "p4")]


def test_comando_atendimento_imprime_e_nao_altera_nada(tmp_path, monkeypatch, capsys):
    registro = _registro_de_exemplo(tmp_path)
    monkeypatch.setattr(ra, "DIRETORIO_PADRAO", registro.parent)
    arquivos = sorted(registro.parent.iterdir())
    antes = {p.name: _hash(p) for p in arquivos}

    assert ra.main(["--atendimento"]) == 0

    saida = capsys.readouterr().out
    assert "leads enviados no registro: 4" in saida
    assert "agente IA: 1" in saida and f"{ra.SEM_ATENDIMENTO}: 1" in saida
    assert "Por nicho:" in saida and "Por pista:" in saida and "Por cidade:" in saida
    assert "Estudio Ficticio p4 [p4]" in saida
    assert {p.name: _hash(p) for p in sorted(registro.parent.iterdir())} == antes, "nada criado nem alterado"
