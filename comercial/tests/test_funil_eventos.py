"""Funil por eventos no registro de abordagens (decisão do diretor,
08/10/2026): cada mudança de Resultado de um lead já registrado vira uma
linha nova, com data; o registro continua só acrescentando, e o funil de
cada lead (envio -> pediu exemplo -> proposta -> cliente) se reconstrói pela
ordem das linhas. Dados fictícios, tudo em tmp_path. Cada regra tem o caso
que grava e o controle que não grava."""

import datetime

import pytest
from openpyxl import Workbook, load_workbook

import registro_abordagens as ra

_FUNIL = {"etapa": ["1º contato", "follow-up 1"],
          "resultado": ["sem resposta", "resposta positiva", "pediu exemplo", "proposta", "cliente"],
          "mensagem": ["enviada como está"], "decidi_nao_enviar": ["sim"]}
_CAB = ["place_id", "nome", "telefone", "Data 1º contato", "Etapa", "Resultado", "Canal", "Ângulo",
        "campanha", "nicho", "pista"]


def _planilha(caminho, linhas, abas=("Geral",)):
    wb = Workbook()
    wb.remove(wb.active)
    for aba in abas:
        ws = wb.create_sheet(aba)
        ws.append(_CAB)
        for l in linhas:
            ws.append([l.get(c) for c in _CAB])
    wb.save(caminho)
    return caminho


def _lead(pid="p1", **over):
    base = {"place_id": pid, "nome": f"Estudio Ficticio {pid}", "telefone": "600000101",
            "Data 1º contato": "2026-10-09", "Etapa": "1º contato", "Canal": "whatsapp",
            "Ângulo": "sem_site", "campanha": "arquitectos-2026-10", "nicho": "Arquitectos", "pista": "direta"}
    base.update(over)
    return base


def _dia(d):
    return datetime.datetime(2026, 10, d, 12, 0, tzinfo=datetime.timezone.utc)


def _registrar(tmp_path, planilha, dia):
    return ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_FUNIL,
                               agora=_dia(dia))


def _linhas_do_registro(tmp_path):
    wb = load_workbook(ra.caminho_registro(tmp_path / "_abordagens"), read_only=True)
    try:
        linhas = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    return [dict(zip(linhas[0], l)) for l in linhas[1:]]


def test_cada_mudanca_de_resultado_vira_uma_linha_nova_com_data(tmp_path):
    r = _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    assert (r["acrescentados"], r["eventos"]) == (1, 0)
    for dia, resultado in ((12, "pediu exemplo"), (15, "proposta"), (20, "cliente")):
        r = _registrar(tmp_path, _planilha(tmp_path / f"{dia}.xlsx", [_lead(Resultado=resultado)]), dia)
        assert (r["acrescentados"], r["eventos"], r["ja_existentes"]) == (0, 1, 0)
    linhas = _linhas_do_registro(tmp_path)
    assert [(l["evento"], l["Resultado"], l["registrado_em"][:10]) for l in linhas] == [
        ("envio", None, "2026-10-09"),
        ("resultado", "pediu exemplo", "2026-10-12"),
        ("resultado", "proposta", "2026-10-15"),
        ("resultado", "cliente", "2026-10-20"),
    ]
    # a linha do evento carrega o lead inteiro: o funil se cruza com nicho, pista e ângulo
    assert {(l["place_id"], l["nicho"], l["pista"], l["angulo"]) for l in linhas} == {
        ("p1", "Arquitectos", "direta", "sem_site")}


def test_funil_por_lead_reconstroi_a_sequencia(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead("p1"), _lead("p2")]), 9)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead("p1", Resultado="pediu exemplo"), _lead("p2")]), 12)
    _registrar(tmp_path, _planilha(tmp_path / "c.xlsx", [_lead("p1", Resultado="proposta"),
                                                          _lead("p2", Resultado="sem resposta")]), 15)
    funil = ra.funil_por_lead(ra.caminho_registro(tmp_path / "_abordagens"))
    assert [(p["evento"], p["resultado"]) for p in funil["p1"]["passos"]] == [
        ("envio", ""), ("resultado", "pediu exemplo"), ("resultado", "proposta")]
    assert [(p["evento"], p["resultado"]) for p in funil["p2"]["passos"]] == [("envio", ""), ("resultado", "sem resposta")]
    resumo = ra.resumo_funil(funil, tuple(_FUNIL["resultado"]))
    assert resumo == {"sem resposta": 1, "resposta positiva": 0, "pediu exemplo": 1, "proposta": 1, "cliente": 0}


def test_mesmo_resultado_de_novo_nao_grava_nada(tmp_path):
    """Controle negativo: rodar o registro outra vez com a mesma planilha não
    cria evento -- só mudança grava."""
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    planilha = _planilha(tmp_path / "b.xlsx", [_lead(Resultado="pediu exemplo")])
    assert _registrar(tmp_path, planilha, 12)["eventos"] == 1
    antes = ra.caminho_registro(tmp_path / "_abordagens").read_bytes()
    r = _registrar(tmp_path, planilha, 13)
    assert (r["acrescentados"], r["eventos"], r["ja_existentes"], r["backup"]) == (0, 0, 1, None)
    assert ra.caminho_registro(tmp_path / "_abordagens").read_bytes() == antes


@pytest.mark.parametrize("resultado", [None, "", "   "])
def test_resultado_vazio_ou_apagado_nao_vira_evento(tmp_path, resultado):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead(Resultado="resposta positiva")]), 9)
    r = _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead(Resultado=resultado)]), 12)
    assert r["eventos"] == 0
    assert len(_linhas_do_registro(tmp_path)) == 1


def test_linhas_antigas_nunca_sao_alteradas(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead("p1"), _lead("p2")]), 9)
    antes = _linhas_do_registro(tmp_path)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead("p1", Resultado="pediu exemplo"), _lead("p2")]), 12)
    depois = _linhas_do_registro(tmp_path)
    assert depois[: len(antes)] == antes
    assert len(depois) == len(antes) + 1


def test_lead_nas_duas_abas_gera_um_evento_so(tmp_path):
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    r = _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead(Resultado="proposta")], abas=("Geral", "Nata")), 12)
    assert r["eventos"] == 1
    assert len(_linhas_do_registro(tmp_path)) == 2


def test_registro_anterior_ao_funil_ganha_a_coluna_e_conta_a_linha_antiga_como_envio(tmp_path):
    """Registro sem a coluna `evento` (antes de 08/10/2026): continua legível,
    a coluna entra no FIM e a linha antiga conta como envio."""
    diretorio = tmp_path / "_abordagens"
    diretorio.mkdir()
    cabecalho_antigo = list(ra.COLUNAS_REGISTRO[:-1])
    wb = Workbook()
    wb.active.append(cabecalho_antigo)
    antiga = {c: "" for c in cabecalho_antigo}
    antiga.update(place_id="p1", nome="Estudio Ficticio p1", registrado_em="2026-10-01T10:00:00Z",
                  enviado_em="2026-10-01", Resultado="resposta positiva")
    wb.active.append([antiga[c] for c in cabecalho_antigo])
    wb.save(ra.caminho_registro(diretorio))

    r = _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead(Resultado="pediu exemplo")]), 12)
    assert r["eventos"] == 1
    linhas = _linhas_do_registro(tmp_path)
    assert list(linhas[0])[-1] == "evento"
    assert linhas[0]["evento"] is None and linhas[1]["evento"] == "resultado"
    funil = ra.funil_por_lead(ra.caminho_registro(diretorio))
    assert [(p["evento"], p["resultado"]) for p in funil["p1"]["passos"]] == [
        ("envio", "resposta positiva"), ("resultado", "pediu exemplo")]


def test_lead_com_eventos_continua_fora_das_proximas_planilhas(tmp_path):
    """`place_ids_registrados` (usado pela planilha de envio para excluir quem
    já foi abordado) segue sendo um conjunto: as linhas de evento não
    duplicam nem somem nenhum lead."""
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead("p1"), _lead("p2")]), 9)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead("p1", Resultado="proposta")]), 12)
    assert ra.place_ids_registrados(ra.caminho_registro(tmp_path / "_abordagens")) == {"p1", "p2"}


def test_comando_funil_mostra_a_sequencia_e_o_resumo(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(ra, "DIRETORIO_PADRAO", tmp_path / "_abordagens")
    monkeypatch.setattr(ra, "carregar_opcoes_funil", lambda *a, **k: _FUNIL)
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead(Resultado="pediu exemplo")]), 12)
    _registrar(tmp_path, _planilha(tmp_path / "c.xlsx", [_lead(Resultado="proposta")]), 15)
    antes = ra.caminho_registro(tmp_path / "_abordagens").read_bytes()

    assert ra.main(["--funil"]) == 0
    saida = capsys.readouterr().out
    assert "Estudio Ficticio p1 [p1]: 2026-10-09 envio -> 2026-10-12 pediu exemplo -> 2026-10-15 proposta" in saida
    assert "Leads no registro: 1" in saida and "pediu exemplo: 1" in saida and "proposta: 1" in saida
    assert "cliente:" not in saida  # etapa sem ninguém não aparece

    assert ra.main(["--funil", "--place-id", "p1"]) == 0
    assert "Leads no registro" not in capsys.readouterr().out
    assert ra.main(["--funil", "--place-id", "p9"]) == 1
    assert ra.caminho_registro(tmp_path / "_abordagens").read_bytes() == antes  # só leitura


def _retrato_da_pasta(pasta):
    """Lista de arquivos (recursiva) com o hash de cada um."""
    import hashlib

    return {
        str(p.relative_to(pasta)): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in sorted(pasta.rglob("*")) if p.is_file()
    }


def test_comando_funil_deixa_a_pasta_do_registro_intacta(tmp_path, monkeypatch, capsys):
    """P2 da verificação (08/10/2026): não basta o registro ter os mesmos
    bytes -- nenhum arquivo novo, apagado ou alterado na pasta (backup,
    temporário, registro novo)."""
    pasta = tmp_path / "_abordagens"
    monkeypatch.setattr(ra, "DIRETORIO_PADRAO", pasta)
    monkeypatch.setattr(ra, "carregar_opcoes_funil", lambda *a, **k: _FUNIL)
    _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead()]), 9)
    _registrar(tmp_path, _planilha(tmp_path / "b.xlsx", [_lead(Resultado="pediu exemplo")]), 12)
    antes = _retrato_da_pasta(pasta)
    assert len(antes) >= 2  # o registro e o backup da segunda gravação: a pasta não está vazia por engano
    for argv in (["--funil"], ["--funil", "--place-id", "p1"], ["--funil", "--place-id", "p9"]):
        ra.main(argv)
    capsys.readouterr()
    assert _retrato_da_pasta(pasta) == antes


def test_data_do_evento_e_a_data_e_hora_locais_do_registro(tmp_path):
    """C4 (diretor, 09/10/2026): `registrado_em` na data e hora LOCAIS do
    computador, sem "Z". Um instante em UTC perto da meia-noite sai convertido
    para o fuso local (na máquina do diretor, Madri)."""
    import datetime as dt

    instante_utc = dt.datetime(2026, 10, 9, 23, 30, tzinfo=dt.timezone.utc)
    ra.registrar_envios(_planilha(tmp_path / "a.xlsx", [_lead()]), diretorio_registro=tmp_path / "_abordagens",
                        opcoes_funil=_FUNIL, agora=instante_utc)
    gravado = _linhas_do_registro(tmp_path)[0]["registrado_em"]
    assert gravado == instante_utc.astimezone().strftime("%Y-%m-%dT%H:%M:%S")
    assert not gravado.endswith("Z")
    assert ra.data_local_do_registro(dt.datetime(2026, 10, 9, 23, 30)) == "2026-10-09T23:30:00"  # sem fuso: já é local


def test_sem_agora_o_registro_usa_o_relogio_local(tmp_path, monkeypatch):
    """Sem `agora`, a hora é a do relógio local (`datetime.now()` sem fuso),
    não UTC: o valor gravado fica a menos de um minuto do relógio local."""
    import datetime as dt

    antes = dt.datetime.now()
    ra.registrar_envios(_planilha(tmp_path / "a.xlsx", [_lead()]), diretorio_registro=tmp_path / "_abordagens",
                        opcoes_funil=_FUNIL)
    gravado = dt.datetime.strptime(_linhas_do_registro(tmp_path)[0]["registrado_em"], "%Y-%m-%dT%H:%M:%S")
    assert abs((gravado - antes).total_seconds()) < 60


def test_lead_novo_nas_duas_abas_entra_uma_vez_so(tmp_path):
    """D3 da verificação (08/10/2026): antes, um lead novo nas duas abas era
    gravado duas vezes como envio ("envio -> envio" no funil)."""
    r = _registrar(tmp_path, _planilha(tmp_path / "a.xlsx", [_lead("p1"), _lead("p2")], abas=("Geral", "Nata")), 9)
    assert r["acrescentados"] == 2
    linhas = _linhas_do_registro(tmp_path)
    assert [l["place_id"] for l in linhas] == ["p1", "p2"]
    funil = ra.funil_por_lead(ra.caminho_registro(tmp_path / "_abordagens"))
    assert [p["evento"] for p in funil["p1"]["passos"]] == ["envio"]


def test_comando_sem_planilha_e_sem_funil_e_erro_de_uso():
    with pytest.raises(SystemExit) as exc:
        ra.main([])
    assert exc.value.code == 2
