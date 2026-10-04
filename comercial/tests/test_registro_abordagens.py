"""Testes de registro_abordagens.py — nenhum envio, nenhuma abertura de
WhatsApp. Toda saída de arquivo em tmp_path, nunca em comercial/_abordagens
real. A planilha do diretor é sempre aberta só leitura -- nunca alterada."""

from datetime import datetime, timezone

import pytest
from openpyxl import Workbook, load_workbook

import registro_abordagens as ra

_OPCOES_FUNIL = {
    "etapa": ["1º contato", "follow-up 1", "follow-up 2"],
    "resultado": ["sem resposta", "resposta positiva", "cliente"],
    "mensagem": ["enviada como está", "editada por mim"],
    "decidi_nao_enviar": ["sim"],
}


def _planilha_com_envios(caminho, linhas_geral=None, linhas_nata=None, cabecalho_extra=()):
    cabecalho = ["place_id", "nome", "telefone", "Data 1º contato", *cabecalho_extra]
    wb = Workbook()
    ws_geral = wb.active
    ws_geral.title = "Geral"
    ws_geral.append(cabecalho)
    for linha in linhas_geral or []:
        ws_geral.append(linha)

    ws_nata = wb.create_sheet("Nata")
    ws_nata.append(cabecalho)
    for linha in linhas_nata or []:
        ws_nata.append(linha)

    wb.save(caminho)
    return caminho


def test_registrar_envios_acrescenta_linhas_com_enviado_em_preenchido(tmp_path):
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-24"], ["p2", "Clínica B", "912345678", ""]],
    )
    resultado = ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens")
    assert resultado["acrescentados"] == 1  # só p1 tem "Data 1º contato" preenchido
    assert resultado["ja_existentes"] == 0

    caminho_reg = ra.caminho_registro(tmp_path / "_abordagens")
    wb = load_workbook(caminho_reg)
    linhas = list(wb.active.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 1
    assert linhas[0][0] == "p1"


def test_registrar_envios_nunca_duplica_place_id(tmp_path):
    diretorio = tmp_path / "_abordagens"
    planilha1 = _planilha_com_envios(tmp_path / "p1.xlsx", linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-24"]])
    ra.registrar_envios(planilha1, diretorio_registro=diretorio)

    # segunda planilha cita o MESMO place_id de novo (ex.: reenvio registrado por engano)
    planilha2 = _planilha_com_envios(tmp_path / "p2.xlsx", linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-25"]])
    resultado2 = ra.registrar_envios(planilha2, diretorio_registro=diretorio)

    assert resultado2["acrescentados"] == 0
    assert resultado2["ja_existentes"] == 1

    wb = load_workbook(ra.caminho_registro(diretorio))
    linhas = list(wb.active.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 1  # continua só uma linha para p1


def test_registrar_envios_le_das_duas_abas_geral_e_nata(tmp_path):
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-24"]],
        linhas_nata=[["p9", "Clínica Nata", "712345678", "2026-09-24"]],
    )
    resultado = ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens")
    assert resultado["acrescentados"] == 2
    wb = load_workbook(ra.caminho_registro(tmp_path / "_abordagens"))
    ids = {linha[0] for linha in wb.active.iter_rows(min_row=2, values_only=True)}
    assert ids == {"p1", "p9"}


def test_registrar_envios_faz_backup_antes_de_gravar(tmp_path):
    diretorio = tmp_path / "_abordagens"
    planilha1 = _planilha_com_envios(tmp_path / "p1.xlsx", linhas_geral=[["p1", "A", "600000101", "2026-09-24"]])
    ra.registrar_envios(planilha1, diretorio_registro=diretorio)

    planilha2 = _planilha_com_envios(tmp_path / "p2.xlsx", linhas_geral=[["p2", "B", "712345678", "2026-09-25"]])
    resultado2 = ra.registrar_envios(planilha2, diretorio_registro=diretorio)

    assert resultado2["backup"] is not None
    assert resultado2["backup"].is_file()


def test_registrar_envios_nunca_altera_a_planilha_do_diretor(tmp_path):
    """Controle: o conteúdo da planilha de origem continua idêntico depois
    de registrar_envios -- só leitura."""
    planilha = _planilha_com_envios(tmp_path / "planilha.xlsx", linhas_geral=[["p1", "A", "600000101", "2026-09-24"]])
    conteudo_antes = planilha.read_bytes()
    ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens")
    assert planilha.read_bytes() == conteudo_antes


def test_registrar_envios_sem_nenhum_enviado_nao_cria_registro(tmp_path):
    planilha = _planilha_com_envios(tmp_path / "planilha.xlsx", linhas_geral=[["p1", "A", "600000101", ""]])
    resultado = ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens")
    assert resultado["acrescentados"] == 0
    assert not ra.caminho_registro(tmp_path / "_abordagens").exists()


def test_registrar_envios_arquivo_de_registro_aberto_da_erro_claro_sem_perder_nada(tmp_path, monkeypatch):
    """Simula o registro estar aberto no Excel (PermissionError no save) --
    nada pode ser perdido: backup feito, planilha do diretor intocada."""
    diretorio = tmp_path / "_abordagens"
    planilha1 = _planilha_com_envios(tmp_path / "p1.xlsx", linhas_geral=[["p1", "A", "600000101", "2026-09-24"]])
    ra.registrar_envios(planilha1, diretorio_registro=diretorio)
    conteudo_registro_antes = ra.caminho_registro(diretorio).read_bytes()

    planilha2 = _planilha_com_envios(tmp_path / "p2.xlsx", linhas_geral=[["p2", "B", "712345678", "2026-09-25"]])

    def _save_bloqueado(self, caminho):
        raise PermissionError("arquivo em uso")

    monkeypatch.setattr("openpyxl.Workbook.save", _save_bloqueado)

    with pytest.raises(ra.PlanilhaAbertaError):
        ra.registrar_envios(planilha2, diretorio_registro=diretorio)

    # nada perdido: registro continua exatamente como estava antes da tentativa
    assert ra.caminho_registro(diretorio).read_bytes() == conteudo_registro_antes
    # a planilha de origem (planilha2), que nem chegou a ser alterada, também intacta
    assert planilha2.is_file()


# --- main() (comando `python ferramentas/registro_abordagens.py <planilha>`) --


def test_main_registra_envios_da_planilha_passada_e_imprime_resumo(tmp_path, monkeypatch, capsys):
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx", linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-24"]]
    )
    monkeypatch.setattr(ra, "caminho_registro", lambda diretorio=None: tmp_path / "_abordagens" / "registro_abordagens.xlsx")

    codigo = ra.main([str(planilha)])

    assert codigo == 0
    saida = capsys.readouterr().out
    assert "1" in saida
    caminho_reg = tmp_path / "_abordagens" / "registro_abordagens.xlsx"
    assert caminho_reg.is_file()


def test_main_exige_o_caminho_da_planilha(capsys):
    """Controle: sem o argumento posicional, argparse recusa (exit != 0) em
    vez de rodar com um caminho inventado."""
    with pytest.raises(SystemExit) as exc:
        ra.main([])
    assert exc.value.code != 0


# --- funil de resultados (decisão do diretor, 24/09/2026, quarta rodada) --


def test_registro_guarda_canal_modelo_e_custo_do_lead(tmp_path):
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_nata=[["p9", "Clínica Nata", "712345678", "2026-09-24", "whatsapp", "claude-sonnet-5", "0.0123"]],
        cabecalho_extra=["canal_sugerido", "modelo", "custo_total_usd"],
    )
    ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)

    wb = load_workbook(ra.caminho_registro(tmp_path / "_abordagens"))
    cabecalho = [c.value for c in wb.active[1]]
    linha = dict(zip(cabecalho, next(wb.active.iter_rows(min_row=2, values_only=True))))
    assert linha["canal"] == "whatsapp"
    assert linha["modelo"] == "claude-sonnet-5"
    assert linha["custo_usd"] == "0.0123"


def test_registro_pega_lead_marcado_decidi_nao_enviar(tmp_path):
    """Prova central da regra nova: um lead que o diretor marcou "Decidi não
    enviar" = "sim" entra no registro mesmo com "Data 1º contato" vazio -- para
    não voltar nas próximas planilhas."""
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["p1", "Clínica A", "600000101", "", "sim", "não quer contato"]],
        cabecalho_extra=["Decidi não enviar", "Motivo de não enviar"],
    )
    resultado = ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)

    assert resultado["acrescentados"] == 1
    assert ra.place_ids_registrados(ra.caminho_registro(tmp_path / "_abordagens")) == {"p1"}


def test_registro_lead_decidi_nao_enviar_vazio_nao_e_pego(tmp_path):
    """Controle negativo: qualquer valor diferente de 'sim' (vazio, 'não',
    etc.) não conta como decisão de não enviar."""
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["p1", "Clínica A", "600000101", "", ""]],
        cabecalho_extra=["Decidi não enviar"],
    )
    resultado = ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)
    assert resultado["acrescentados"] == 0


def test_registro_e_append_only_nunca_altera_linha_existente(tmp_path):
    """Prova direta de "só acrescenta, nunca altera": depois de registrar
    p1, uma segunda rodada com uma planilha que traz p1 de novo (com dados
    diferentes) não muda a linha de p1 já gravada -- só ignora (conta em
    ja_existentes)."""
    diretorio = tmp_path / "_abordagens"
    planilha1 = _planilha_com_envios(tmp_path / "p1.xlsx", linhas_geral=[["p1", "Nome Original", "600000101", "2026-09-24"]])
    ra.registrar_envios(planilha1, diretorio_registro=diretorio, opcoes_funil=_OPCOES_FUNIL)

    caminho_reg = ra.caminho_registro(diretorio)
    linha_antes = list(load_workbook(caminho_reg).active.iter_rows(min_row=2, values_only=True))[0]

    planilha2 = _planilha_com_envios(tmp_path / "p2.xlsx", linhas_geral=[["p1", "Nome Editado Por Engano", "999999999", "2026-09-25"]])
    resultado2 = ra.registrar_envios(planilha2, diretorio_registro=diretorio, opcoes_funil=_OPCOES_FUNIL)

    assert resultado2["acrescentados"] == 0
    assert resultado2["ja_existentes"] == 1
    linha_depois = list(load_workbook(caminho_reg).active.iter_rows(min_row=2, values_only=True))[0]
    assert linha_depois == linha_antes  # nada mudou -- o registro nunca altera linha existente


def test_registro_aplica_validacao_de_lista_nas_colunas_do_funil(tmp_path):
    planilha = _planilha_com_envios(tmp_path / "planilha.xlsx", linhas_geral=[["p1", "A", "600000101", "2026-09-24"]])
    ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)

    wb = load_workbook(ra.caminho_registro(tmp_path / "_abordagens"))
    validacoes = wb.active.data_validations.dataValidation
    assert len(validacoes) >= 4  # Etapa, Resultado, Enviada como, Decidi não enviar


def test_registro_valida_a_coluna_enviada_como_no_lugar_de_mensagem(tmp_path):
    """Defeito relatado pela supervisão, 24/09/2026: a coluna "Mensagem" foi
    renomeada para "Enviada como" -- a validação de lista tem que seguir o
    nome novo, não silenciosamente parar de aplicar."""
    planilha = _planilha_com_envios(tmp_path / "planilha.xlsx", linhas_geral=[["p1", "A", "600000101", "2026-09-24"]])
    ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)

    wb = load_workbook(ra.caminho_registro(tmp_path / "_abordagens"))
    cabecalho = [c.value for c in wb.active[1]]
    idx_enviada_como = cabecalho.index("Enviada como")
    letra = wb.active.cell(row=1, column=idx_enviada_como + 1).column_letter
    assert any(letra in dv.sqref.__str__() for dv in wb.active.data_validations.dataValidation)


def test_carregar_opcoes_funil_le_o_arquivo_real():
    config = ra.carregar_opcoes_funil()
    for chave in ("etapa", "resultado", "mensagem", "decidi_nao_enviar"):
        assert chave in config


def test_carregar_opcoes_funil_arquivo_ausente_levanta_erro_claro(tmp_path):
    with pytest.raises(ra.ConfigFunilInvalidaError, match="não encontrada"):
        ra.carregar_opcoes_funil(tmp_path / "nao-existe.json")


def test_aplicar_validacao_lista_sem_opcoes_nao_aplica_nada(tmp_path):
    """Controle negativo: opções vazias -- nenhuma validação criada."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Etapa"])
    ws.append(["1º contato"])
    ra.aplicar_validacao_lista(ws, ("Etapa",), "Etapa", [], ultima_linha=2)
    assert len(ws.data_validations.dataValidation) == 0


# --- place_id como identificador padrão (defeito grave, 25/09/2026) --------
#
# A aba Nata não tinha a coluna place_id: _linhas_enviadas pulava a aba
# inteira em silêncio (nenhum envio dela era registrado, os leads voltavam
# nas próximas planilhas). Agora: aba com envio pendente e sem a coluna, ou
# com place_id vazio numa linha pendente, levanta erro claro -- nunca pula
# em silêncio, nunca cai para nome ou telefone.


def _planilha_sem_coluna_place_id(caminho, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Nata"
    ws.append(["nome", "telefone", "Data 1º contato"])  # sem place_id
    for linha in linhas:
        ws.append(linha)
    wb.save(caminho)
    return caminho


def test_registrar_envios_aba_sem_coluna_place_id_levanta_erro_claro_e_nao_registra_nada(tmp_path):
    planilha = _planilha_sem_coluna_place_id(
        tmp_path / "planilha.xlsx", [["Clínica A", "600000101", "2026-09-25"]]
    )
    with pytest.raises(ra.RegistroSemPlaceIdError, match='"Nata"'):
        ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)
    assert not ra.caminho_registro(tmp_path / "_abordagens").exists()


def test_registrar_envios_linha_com_place_id_vazio_levanta_erro_claro_e_nao_registra_nada(tmp_path):
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-25"], ["", "Clínica B", "912345678", "2026-09-25"]],
    )
    with pytest.raises(ra.RegistroSemPlaceIdError, match='"Geral"'):
        ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)
    assert not ra.caminho_registro(tmp_path / "_abordagens").exists()


def test_registrar_envios_erro_de_uma_aba_impede_registro_da_outra_tambem(tmp_path):
    """Prova central: "nunca registra nada pela metade" -- mesmo com a aba
    Geral totalmente válida, um problema na Nata bloqueia o registro
    inteiro."""
    wb = Workbook()
    ws_geral = wb.active
    ws_geral.title = "Geral"
    ws_geral.append(["place_id", "nome", "telefone", "Data 1º contato"])
    ws_geral.append(["p1", "Clínica A", "600000101", "2026-09-25"])
    ws_nata = wb.create_sheet("Nata")
    ws_nata.append(["nome", "telefone", "Data 1º contato"])  # sem place_id
    ws_nata.append(["Clínica B", "712345678", "2026-09-25"])
    caminho = tmp_path / "planilha.xlsx"
    wb.save(caminho)

    with pytest.raises(ra.RegistroSemPlaceIdError):
        ra.registrar_envios(caminho, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)
    assert not ra.caminho_registro(tmp_path / "_abordagens").exists()


def test_registrar_envios_decidi_nao_enviar_sem_place_id_tambem_levanta_erro(tmp_path):
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["", "Clínica A", "600000101", "", "sim"]],
        cabecalho_extra=["Decidi não enviar"],
    )
    with pytest.raises(ra.RegistroSemPlaceIdError):
        ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)


# --- ponta a ponta: gerar planilha real -> preencher -> registrar -> excluir


def test_ponta_a_ponta_gerar_preencher_registrar_e_excluir_na_proxima_planilha(tmp_path, monkeypatch):
    """Teste pedido explicitamente pela supervisão, 25/09/2026: gera a
    planilha real (planilha_envio.gerar_planilha), preenche "Data 1º
    contato" numa linha de CADA aba (Geral e Nata), registra, confirma que
    as duas aparecem no registro, e que a PRÓXIMA planilha exclui as duas."""
    import contrato_loader
    import planilha_envio as pe
    from contrato_loader import LeadQualificado
    from types import SimpleNamespace

    def _lead_nata(place_id):
        return LeadQualificado({
            "place_id": place_id,
            "classe_site": "proprio",
            "identidade": {"nome": "Clínica Nata", "nicho": "Fisio", "cidade": None, "endereco": None, "google_maps_url": None},
            "contato": {
                "telefone": {"valor": "712345678", "estado": "CONFIRMADO_PRESENTE"},
                "email": {"valor": None, "estado": "NAO_VERIFICADO"},
                "whatsapp_apto": True,
            },
            "presenca_digital": {"website_url": {"valor": None, "estado": "NAO_VERIFICADO"}},
            "analise_tecnica_site": {
                "status": "ok", "error_type": None, "avaliado_em": None,
                "telefone_na_pagina": {"valor": "texto", "estado": "CONFIRMADO_PRESENTE"},
            },
            "reputacao": {
                "nota_google": {"valor": None, "estado": "NAO_VERIFICADO"},
                "review_count": {"valor": None, "estado": "NAO_VERIFICADO"},
            },
            "problema_vendavel": [],
            "psi": {"estado": "NAO_VERIFICADO", "lcp_ms": None, "motivo": None, "medido_em": None, "rodadas": 1},
            "texto_site": None,
        })

    import csv

    linha_csv = {
        "pista": "direta", "motivo": "sem site", "nome": "Clínica Geral", "cidade": "",
        "telefone": "600000101", "email": "", "instagram": "", "site": "",
        "classe_site": "sem_site", "avaliacoes": "27", "nota": "4.8",
        "google_maps_url": "https://maps.google.com/?cid=1", "place_id": "p-geral",
    }
    caminho_csv = tmp_path / "qualificador_20260925-090000.csv"
    with caminho_csv.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(linha_csv.keys()), delimiter=";")
        w.writeheader()
        w.writerow(linha_csv)

    lote = SimpleNamespace(nata=[_lead_nata("p-nata")], candidatos_triagem=[], descartados=[])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: lote)

    diretorio_registro = tmp_path / "_abordagens"
    saida_dir = tmp_path / "_planilhas"

    # 1) gera a planilha real
    # consolidada_dir isolado em tmp_path -- sem isto, gerar_planilha cairia no
    # COMERCIAL_PLANILHAS_DIR real (comercial/.env) e escreveria fixture de
    # teste ("p-geral"/"p-nata") na planilha consolidada de verdade do diretor
    # (defeito real, achado e limpo manualmente em 01/10/2026).
    resultado1 = pe.gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=saida_dir, caminho_registro=diretorio_registro / "registro_abordagens.xlsx",
        agora=datetime(2026, 9, 25, 9, 0, 0, tzinfo=timezone.utc), diretorio_capturas=tmp_path / "_capturas_inexistente",
        consolidada_dir=tmp_path / "_consolidada_teste",
    )
    assert resultado1["excluidos_geral"] == 0
    assert resultado1["excluidos_nata"] == 0

    # 2) preenche "Data 1º contato" numa linha de cada aba
    wb = load_workbook(resultado1["caminho"])
    for nome_aba, place_id_alvo in (("Geral", "p-geral"), ("Nata", "p-nata")):
        ws = wb[nome_aba]
        cabecalho = [c.value for c in ws[1]]
        idx_place_id = cabecalho.index("place_id") + 1
        idx_data = cabecalho.index("Data 1º contato") + 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=idx_place_id).value == place_id_alvo:
                ws.cell(row=r, column=idx_data).value = "2026-09-25"
    wb.save(resultado1["caminho"])

    # 3) registra
    resultado_registro = ra.registrar_envios(resultado1["caminho"], diretorio_registro=diretorio_registro)
    assert resultado_registro["acrescentados"] == 2

    registrados = ra.place_ids_registrados(ra.caminho_registro(diretorio_registro))
    assert registrados == {"p-geral", "p-nata"}

    # 4) a próxima planilha exclui as duas
    resultado2 = pe.gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=saida_dir, caminho_registro=diretorio_registro / "registro_abordagens.xlsx",
        agora=datetime(2026, 9, 25, 9, 30, 0, tzinfo=timezone.utc), diretorio_capturas=tmp_path / "_capturas_inexistente",
        consolidada_dir=tmp_path / "_consolidada_teste",
    )
    assert resultado2["excluidos_geral"] == 1
    assert resultado2["excluidos_nata"] == 1


# --- teste A/B da abertura (decisão do diretor, 27/09/2026) -------------------------


def _linha_do_registro(tmp_path):
    wb = load_workbook(ra.caminho_registro(tmp_path / "_abordagens"))
    cabecalho = [c.value for c in wb.active[1]]
    return cabecalho, dict(zip(cabecalho, next(wb.active.iter_rows(min_row=2, values_only=True))))


def test_registro_guarda_canal_angulo_e_variante_da_planilha_atual(tmp_path):
    """Defeito achado em 27/09/2026: o registro só lia "canal_sugerido"
    (planilha da Etapa E2); a planilha atual tem "Canal" -- o canal chegava
    vazio, e ângulo/variante nem eram lidos. Sem eles o A/B não se separa."""
    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-27", "whatsapp", "sem_site", "B"]],
        cabecalho_extra=["Canal", "Ângulo", "Variante"],
    )
    ra.registrar_envios(planilha, diretorio_registro=tmp_path / "_abordagens", opcoes_funil=_OPCOES_FUNIL)
    _, linha = _linha_do_registro(tmp_path)
    assert linha["canal"] == "whatsapp"
    assert linha["angulo"] == "sem_site"
    assert linha["variante"] == "B"


def test_registro_antigo_ganha_colunas_novas_no_fim_sem_mexer_nas_antigas(tmp_path):
    """Registro criado antes de 27/09 (cabeçalho sem angulo/variante): as
    colunas novas entram no FIM; a linha antiga fica intacta e a nova é
    gravada por nome de coluna, não por posição."""
    diretorio = tmp_path / "_abordagens"
    diretorio.mkdir()
    antigas = [c for c in ra.COLUNAS_REGISTRO if c not in ("angulo", "variante")]
    wb = Workbook()
    ws = wb.active
    ws.append(antigas)
    linha_antiga = ["p0", "Clínica Vieja", "600000000", "Geral", "2026-09-25", "email", "", "", "2026-09-25T10:00:00Z",
                    "1º contato", "sem resposta", "", "", "", ""]
    ws.append(linha_antiga)
    wb.save(ra.caminho_registro(diretorio))

    planilha = _planilha_com_envios(
        tmp_path / "planilha.xlsx",
        linhas_geral=[["p1", "Clínica A", "600000101", "2026-09-27", "whatsapp", "portal", "A"]],
        cabecalho_extra=["Canal", "Ângulo", "Variante"],
    )
    ra.registrar_envios(planilha, diretorio_registro=diretorio, opcoes_funil=_OPCOES_FUNIL)

    wb = load_workbook(ra.caminho_registro(diretorio))
    cabecalho = [c.value for c in wb.active[1]]
    assert cabecalho == antigas + ["angulo", "variante"]
    linhas = list(wb.active.iter_rows(min_row=2, values_only=True))
    assert list(linhas[0][: len(antigas)]) == [v if v != "" else None for v in linha_antiga]
    nova = dict(zip(cabecalho, linhas[1]))
    assert nova["place_id"] == "p1"
    assert nova["canal"] == "whatsapp"
    assert nova["angulo"] == "portal"
    assert nova["variante"] == "A"
