"""Testes de planilha_envio.py — nenhum envio, nenhuma abertura de WhatsApp,
nenhuma chamada de LLM. Toda saída de arquivo em tmp_path, nunca em
comercial/_planilhas real.

MVP de mensagem por modelo fixo (mudança de rumo, decisão do diretor,
24/09/2026): este arquivo foi reescrito quase por inteiro -- a Etapa E2
(chamada de LLM, `resultados.json` de uma rodada) saiu do fluxo da planilha,
substituída por `angulo_mensagem.py`. Ver o RETORNO da tarefa para a lista
dos testes removidos/renomeados nesta rodada."""

import json
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

import angulo_mensagem as am
import planilha_envio as pe
from contrato_loader import LeadQualificado
from planilha_envio import (
    COLUNAS_GERAL,
    COLUNAS_NATA,
    LIMITE_CELULA_XLSX,
    STATUS_SEM_MENSAGEM,
    LinkInconsistenteError,
    PlaceIdAusenteError,
    PlanilhaJaExisteError,
    determinar_canal,
    gerar_planilha,
    ler_csv_humano,
    link_captura,
    montar_contexto_ia_geral,
    montar_contexto_ia_nata,
    montar_linhas_geral,
    montar_linhas_nata,
    place_ids_ja_abordados,
    profissional_detectado,
    verificar_consistencia_links,
    _sem_caracteres_ilegais,
    _validar_place_ids,
    _verificar_links_aba,
)

_ASSUNTOS_EMAIL = {"direta": "Una observación sobre su perfil en Google", "nata": "Una observación sobre su web"}
_OPCOES_FUNIL = {
    "etapa": ["1º contato", "follow-up 1", "follow-up 2"],
    "resultado": ["sem resposta", "resposta positiva", "resposta negativa", "cliente"],
    "mensagem": ["enviada como está", "editada por mim"],
    "decidi_nao_enviar": ["sim"],
}
_REGRAS_ANGULO = {
    "poucas_avaliacoes": {
        "nota_minima": 4.0, "avaliacoes_minimo": 1, "avaliacoes_maximo": 30,
        "multiplicador_concorrente_minimo": 10,
    },
    "lentidao": {"lcp_minimo_ms": 10000, "rodadas_minimas_para_citar_numero": 2},
}
def _m(fato, cta, consequencia=None):
    return {"fato": fato, "consequencia": consequencia, "cta": cta}


# Blocos fato -> consequência -> CTA (decisão do diretor, 27/09/2026) --
# fixture própria, não a copy real.
_MENSAGENS_ANGULO = {
    "saudacao": "Hola, buenas.",
    "contato": _m("entré en su web desde el móvil.", "¿se la envío?"),
    "poucas_avaliacoes_5_30": _m("un {nota} con {n} reseñas.", "¿le cuento?", "Puede encontrarse con otras."),
    "poucas_avaliacoes_5_30_singular": _m("un {nota} con {n} reseñas.", "¿le cuento?", "Puede encontrarse con otra."),
    "poucas_avaliacoes_1_4": _m("todavía tiene solo {n} {reseñas_palavra}.", "¿le cuento?", "Puede encontrarse con otras."),
    "poucas_avaliacoes_1_4_singular": _m("todavía tiene solo {n} {reseñas_palavra}.", "¿le cuento?", "Puede encontrarse con otra."),
    "lentidao": _m("estuve unos {s} segundos esperando.", "¿se lo envío?"),
    "lentidao_sem_numero": _m("estuve bastante rato esperando.", "¿se lo envío?"),
    "sem_site": _m("vi su ficha de Google, un {nota} con {n} reseñas, enhorabuena.", "¿Le llegan más por Google?"),
    "sem_site_sem_reputacao": _m("vi su ficha de Google.", "¿Le llegan más por Google?"),
    "portal": _m("vi su ficha, {nota} con {n} reseñas, y que lleva a {portal}.", "¿Les llegan muchos por {portal}?"),
    "portal_sem_reputacao": _m("vi su ficha y que lleva a {portal}.", "¿Les llegan muchos por {portal}?"),
    "rede_social": _m("vi su ficha, {nota} con {n} reseñas, y que lleva a su perfil de {rede}.", "¿Les escriben por {rede}?"),
    "rede_social_sem_reputacao": _m("vi su ficha y que lleva a su perfil de {rede}.", "¿Les escriben por {rede}?"),
    "construtor": _m("vi su ficha, {nota} con {n} reseñas, y que usa {constructor}.", "¿La montaron ustedes mismos?"),
    "construtor_sem_reputacao": _m("vi su ficha y que usa {constructor}.", "¿La montaron ustedes mismos?"),
    "nomes_portal": {"doctoralia.es": "Doctoralia"},
    "nomes_rede": {"instagram.com": "Instagram"},
    "nomes_construtor": {"wordpress.com": "WordPress"},
}
# Apresentação vazia por padrão: os testes de linha olham o corpo, não a
# abertura (a abertura A/B foi aposentada em 29/09/2026, D12 Fase 5.1).
_APRESENTACAO = ""
_CONFIG_VALIDACAO = {
    "max_palavras_mensagem_1": 80,
    "simbolos_moeda": ["€"],
    "palavras_preco": ["precio", "tarifa"],
    "marcas_informais_es": ["tú", "tu", "te", "tienes", "puedes"],
    "saudacao_fixa": "Hola, buenas",
    "termos_proibidos": ["está caída", "está perdiendo", "se van a otra"],
}


def _kwargs_comuns(diretorio_capturas):
    return dict(
        apresentacao=_APRESENTACAO, config_validacao=_CONFIG_VALIDACAO, assuntos_email=_ASSUNTOS_EMAIL,
        mensagens_angulo=_MENSAGENS_ANGULO, diretorio_capturas=diretorio_capturas,
    )


def _linha_csv(**over):
    base = {
        "pista": "direta", "motivo": "sem site", "nome": "Clínica Ejemplo", "cidade": "",
        "telefone": "600000101", "email": "", "instagram": "", "site": "",
        "classe_site": "sem_site", "avaliacoes": "27", "nota": "4.8",
        "google_maps_url": "https://maps.google.com/?cid=1", "place_id": "p1",
    }
    base.update(over)
    return base


def _lead_nata(
    place_id="p1", nicho="Fisioterapia", nome="Clínica Ejemplo", telefone="600000101", email=None,
    telefone_na_pagina=None, nota_estado="NAO_VERIFICADO", nota_valor=None,
    avaliacoes_estado="NAO_VERIFICADO", avaliacoes_valor=None, lcp_ms=None, psi_estado="NAO_VERIFICADO",
    rodadas=2, texto_site=None, fora_do_ar=False, classe_site="proprio",
):
    problemas = []
    if fora_do_ar:
        problemas.append({"tipo": "fora_do_ar", "evidencia": "x", "medicao": "site_renderizado", "coletado_em": None})
    return LeadQualificado({
        "place_id": place_id,
        "classe_site": classe_site,
        "identidade": {
            "nome": nome, "nicho": nicho, "cidade": None, "endereco": None,
            "google_maps_url": "https://maps.google.com/?cid=1",
        },
        "contato": {
            "telefone": {"valor": telefone, "estado": "CONFIRMADO_PRESENTE" if telefone else "NAO_VERIFICADO"},
            "email": {"valor": email, "estado": "CONFIRMADO_PRESENTE" if email else "NAO_VERIFICADO"},
            "whatsapp_apto": True,
        },
        "presenca_digital": {"website_url": {"valor": "https://clinica.es", "estado": "CONFIRMADO_PRESENTE"}},
        "analise_tecnica_site": {
            "status": "ok", "error_type": None, "avaliado_em": None,
            "telefone_na_pagina": {
                "valor": telefone_na_pagina, "estado": "CONFIRMADO_PRESENTE" if telefone_na_pagina else "NAO_VERIFICADO",
            },
        },
        "reputacao": {
            "nota_google": {"valor": nota_valor, "estado": nota_estado},
            "review_count": {"valor": avaliacoes_valor, "estado": avaliacoes_estado},
        },
        "problema_vendavel": problemas,
        "psi": {"estado": psi_estado, "lcp_ms": lcp_ms, "motivo": None, "medido_em": None, "rodadas": rodadas},
        "texto_site": (
            {"estado": "CONFIRMADO_PRESENTE", "texto": texto_site, "chars_originais": len(texto_site),
             "truncado": False, "dispositivo": "mobile", "coletado_em": None}
            if texto_site else None
        ),
    })


def _lote_falso(nata=(), candidatos_triagem=(), descartados=()):
    return SimpleNamespace(nata=list(nata), candidatos_triagem=list(candidatos_triagem), descartados=list(descartados))


# --- determinar_canal (decisão do diretor, 24/09/2026, quinta rodada) -------


def test_canal_whatsapp_quando_celular():
    assert determinar_canal("600000101", "x@y.es") == "whatsapp"


def test_canal_email_quando_sem_celular_com_email():
    assert determinar_canal("", "x@y.es") == "email"
    assert determinar_canal("912345678", "x@y.es") == "email"  # fixo não conta como celular


def test_canal_para_depois_quando_nenhum():
    assert determinar_canal("", "") == "para depois"
    assert determinar_canal("912345678", "") == "para depois"


def test_canal_celular_tem_prioridade_mesmo_com_email():
    assert determinar_canal("600000101", "x@y.es") == "whatsapp"


# --- profissional_detectado (decisão do diretor, 24/09/2026, quinta rodada) -


@pytest.mark.parametrize("nome", ["Dr. Juan Pérez", "Dra. Ana", "Doctor Juan", "Doctora Ana", "dr juan", "DRA. X"])
def test_profissional_detectado_sim(nome):
    assert profissional_detectado(nome) == "sim"


@pytest.mark.parametrize("nome", ["Clínica Ejemplo", "Centro Médico Doctor Bien", None, ""])
def test_profissional_detectado_nao(nome):
    """Controle negativo: só conta quando o nome COMEÇA com o título --
    "Centro Médico Doctor Bien" tem "Doctor" no meio, não no início."""
    assert profissional_detectado(nome) == "não"


def test_profissional_detectado_nunca_extrai_o_nome():
    """A função só devolve sim/não -- nunca um pedaço do nome."""
    resultado = profissional_detectado("Dra. María López")
    assert resultado in ("sim", "não")
    assert "María" not in resultado and "López" not in resultado


# --- _sem_caracteres_ilegais (defeito real, 25/09/2026) --------------------


def test_sem_caracteres_ilegais_remove_caractere_de_controle():
    assert _sem_caracteres_ilegais("Texto com \x0b caractere de controle") == "Texto com  caractere de controle"
    assert _sem_caracteres_ilegais("Outro \x1f caso") == "Outro  caso"


def test_sem_caracteres_ilegais_texto_normal_fica_igual():
    """Controle negativo: texto sem caractere ilegal não é alterado -- só
    remove o que o Excel de fato recusa."""
    assert _sem_caracteres_ilegais("Ofrecemos fisioterapia deportiva.") == "Ofrecemos fisioterapia deportiva."


def test_sem_caracteres_ilegais_valor_nao_string_passa_intacto():
    assert _sem_caracteres_ilegais(42) == 42
    assert _sem_caracteres_ilegais(None) is None


def test_gerar_planilha_com_caractere_de_controle_no_texto_site_nao_quebra(tmp_path, monkeypatch):
    """Prova central do defeito real: texto_site com \\x0b (caractere de
    controle real achado em produção, 25/09/2026) não pode impedir a
    planilha de ser gerada -- antes, IllegalCharacterError travava
    openpyxl.Workbook.save e NENHUMA linha era gravada."""
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv()])
    lead_com_caractere_ilegal = _lead_nata(texto_site="Ofrecemos fisioterapia\x0bdeportiva y \x1frehabilitación.")
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso(nata=[lead_com_caractere_ilegal]))

    resultado = gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=tmp_path / "_planilhas",
        agora=datetime(2026, 9, 25, 9, 0, 0, tzinfo=timezone.utc), **_kwargs_gerar_planilha(tmp_path),
    )

    wb = load_workbook(resultado["caminho"])
    cab_nata = [c.value for c in wb["Nata"][1]]
    idx_texto_site = cab_nata.index("texto_site") + 1
    valor = wb["Nata"].cell(row=2, column=idx_texto_site).value
    assert "\x0b" not in valor
    assert "\x1f" not in valor
    assert "fisioterapia" in valor and "rehabilitación" in valor


# --- place_id como identificador padrão (defeito grave, 25/09/2026) --------


def test_colunas_nata_tem_place_id():
    """A aba Nata NÃO tinha essa coluna -- registro_abordagens pulava a aba
    inteira em silêncio, e nenhum envio dela era registrado."""
    assert "place_id" in COLUNAS_NATA


def test_colunas_geral_tem_place_id():
    assert "place_id" in COLUNAS_GERAL


def test_montar_linhas_nata_preenche_place_id(tmp_path):
    lead = _lead_nata(place_id="p-real")
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    assert linhas[0]["place_id"] == "p-real"


def test_validar_place_ids_passa_quando_todas_as_linhas_tem(tmp_path):
    _validar_place_ids("Geral", [{"place_id": "p1"}, {"place_id": "p2"}])  # não levanta


def test_validar_place_ids_levanta_erro_claro_com_aba_e_linhas(tmp_path):
    with pytest.raises(PlaceIdAusenteError, match='"Nata"'):
        _validar_place_ids("Nata", [{"place_id": "p1"}, {"place_id": ""}, {"place_id": None}])


def test_gerar_planilha_recusa_gravar_linha_sem_place_id(tmp_path, monkeypatch):
    """Prova central: gerar_planilha nunca grava uma planilha incompleta em
    silêncio -- levanta PlaceIdAusenteError e nenhum arquivo é criado."""
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv(place_id="")])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso())

    saida_dir = tmp_path / "_planilhas"
    with pytest.raises(PlaceIdAusenteError, match='"Geral"'):
        gerar_planilha(
            caminho_csv_humano=caminho_csv, saida_dir=saida_dir,
            agora=datetime(2026, 9, 25, 10, 0, 0, tzinfo=timezone.utc), **_kwargs_gerar_planilha(tmp_path),
        )
    assert not any(saida_dir.glob("*.xlsx"))


def test_gerar_planilha_esconde_a_coluna_place_id(tmp_path, monkeypatch):
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv()])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso(nata=[_lead_nata()]))

    resultado = gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=tmp_path / "_planilhas",
        agora=datetime(2026, 9, 25, 10, 30, 0, tzinfo=timezone.utc), **_kwargs_gerar_planilha(tmp_path),
    )

    wb = load_workbook(resultado["caminho"])
    for nome_aba, colunas in (("Geral", COLUNAS_GERAL), ("Nata", COLUNAS_NATA)):
        ws = wb[nome_aba]
        idx = colunas.index("place_id") + 1
        letra = ws.cell(row=1, column=idx).column_letter
        assert ws.column_dimensions[letra].hidden is True


# --- link_captura (decisão do diretor, 24/09/2026, quinta rodada) ----------


def test_link_captura_arquivo_existente(tmp_path):
    (tmp_path / "p1_mobile.png").write_bytes(b"fake-png")
    valor, aviso = link_captura("p1", "proprio", tmp_path)
    assert valor.startswith("file:")
    assert valor.endswith("p1_mobile.png")
    assert aviso == ""


def test_link_captura_arquivo_ausente(tmp_path):
    valor, aviso = link_captura("p-sem-captura", "proprio", tmp_path)
    assert valor == ""
    assert "não encontrada" in aviso


def test_link_captura_sem_place_id(tmp_path):
    valor, aviso = link_captura(None, "proprio", tmp_path)
    assert valor == ""
    assert aviso != ""


def test_link_captura_classe_site_nao_proprio_fica_vazio_sem_aviso(tmp_path):
    """Defeito relatado pela supervisão, 24/09/2026: captura só é esperada
    para lead com site próprio -- fora disso (Direta: sem_site/portal/...),
    célula vazia e SEM aviso, mesmo sem o arquivo existir."""
    valor, aviso = link_captura("p1", "sem_site", tmp_path)
    assert valor == ""
    assert aviso == ""
    valor, aviso = link_captura(None, "portal", tmp_path)
    assert valor == ""
    assert aviso == ""


# --- montar_linhas_geral: ângulo sem_site / portal --------------------------


def test_linha_geral_sem_site_com_nota_confirmada_ganha_mensagem(tmp_path):
    linhas = montar_linhas_geral([_linha_csv()], **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Ângulo"] == "sem_site"
    assert linha["Mensagem sugerida"].startswith("Hola, buenas.")
    assert linha["Aviso"] == ""  # classe_site != "proprio" -- sem captura esperada, sem aviso


def test_linha_geral_portal_pela_classe_site(tmp_path):
    """Prova central: 'portal' é escolhido pela classe_site (QUALIFICADOR),
    não pelo domínio -- Top Doctors (que não é Doctoralia) também ganha
    mensagem agora."""
    linhas = montar_linhas_geral(
        [_linha_csv(classe_site="portal", site="https://www.topdoctors.es/x")], **_kwargs_comuns(tmp_path)
    )
    assert linhas[0]["Ângulo"] == "portal"
    assert linhas[0]["Mensagem sugerida"] != ""


def test_linha_geral_classe_proprio_fica_sem_angulo_sem_falha_de_mensagem(tmp_path):
    """Controle negativo central: com as 5 classes de "sem site próprio"
    cobertas (sem_site, superficie_google, portal, rede_social,
    construtor), só "proprio" fica sem_angulo -- sem_angulo não é falha de
    mensagem (o Aviso de captura ausente é outra checagem, testada à parte
    em test_link_captura_*)."""
    linhas = montar_linhas_geral(
        [_linha_csv(classe_site="proprio", site="https://clinica.es")], **_kwargs_comuns(tmp_path)
    )
    linha = linhas[0]
    assert linha["Ângulo"] == am.SEM_ANGULO
    assert linha["Mensagem sugerida"] == ""


def test_linha_geral_rede_social_ganha_mensagem(tmp_path):
    linhas = montar_linhas_geral(
        [_linha_csv(classe_site="rede_social", site="https://www.instagram.com/x", nota="4.8", avaliacoes="10")],
        **_kwargs_comuns(tmp_path),
    )
    linha = linhas[0]
    assert linha["Ângulo"] == "rede_social"
    assert linha["Mensagem sugerida"] != ""


def test_linha_geral_construtor_ganha_mensagem(tmp_path):
    linhas = montar_linhas_geral(
        [_linha_csv(classe_site="construtor", site="https://clinica.negocio.site", nota="4.8", avaliacoes="10")],
        **_kwargs_comuns(tmp_path),
    )
    linha = linhas[0]
    assert linha["Ângulo"] == "construtor"
    assert linha["Mensagem sugerida"] != ""


def test_linha_geral_sem_nota_confirmada_usa_variante_sem_reputacao_sem_aviso(tmp_path):
    """Defeito relatado pela supervisão, 25/09/2026: leads sem nota/
    avaliações confirmadas ficavam sem mensagem, com Aviso "nota/avaliações
    não confirmada" -- agora usam a variante sem reputação e saem com
    mensagem, sem esse aviso."""
    linhas = montar_linhas_geral([_linha_csv(nota="", avaliacoes="")], **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Ângulo"] == "sem_site"
    assert linha["Mensagem sugerida"] != ""
    assert "reseñas" not in linha["Mensagem sugerida"]
    assert linha["Aviso"] == ""


def test_linha_geral_zero_avaliacoes_usa_variante_sem_reputacao(tmp_path):
    linhas = montar_linhas_geral([_linha_csv(nota="4.8", avaliacoes="0")], **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Mensagem sugerida"] != ""
    assert "reseñas" not in linha["Mensagem sugerida"]
    assert linha["Aviso"] == ""


def test_linha_geral_portal_sem_reputacao_confirmada_usa_variante(tmp_path):
    linhas = montar_linhas_geral(
        [_linha_csv(classe_site="portal", site="https://www.topdoctors.es/x", nota="", avaliacoes="")],
        **_kwargs_comuns(tmp_path),
    )
    linha = linhas[0]
    assert linha["Ângulo"] == "portal"
    assert linha["Mensagem sugerida"] != ""
    assert "reseñas" not in linha["Mensagem sugerida"]
    assert linha["Aviso"] == ""


def test_linha_geral_com_captura_e_mensagem_valida_fica_sem_nenhum_aviso(tmp_path):
    (tmp_path / "p1_mobile.png").write_bytes(b"x")
    linhas = montar_linhas_geral([_linha_csv(place_id="p1")], **_kwargs_comuns(tmp_path))
    assert linhas[0]["Aviso"] == ""


def test_linha_geral_canal_whatsapp_preenche_whatsapp_com_mensagem_nao_abrir_email(tmp_path):
    """Coluna única "WhatsApp" (decisão do diretor, 25/09/2026): com celular
    e mensagem válida, o link já vem com a mensagem -- "Abrir conversa" saiu."""
    linhas = montar_linhas_geral([_linha_csv(telefone="600000101", email="x@y.es")], **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Canal"] == "whatsapp"
    assert linha["Mensagem sugerida"] != ""
    assert linha["WhatsApp"] != ""
    assert "text=" in linha["WhatsApp"]  # link com a mensagem, não conversa vazia
    assert linha["Abrir e-mail"] == ""


def test_linha_geral_canal_whatsapp_sem_mensagem_valida_usa_conversa_vazia(tmp_path):
    """Controle negativo central: celular presente mas SEM mensagem válida
    (ex.: sem_angulo) -- a coluna "WhatsApp" não fica vazia, é o link de
    conversa vazia (nunca célula em branco quando há celular)."""
    linhas = montar_linhas_geral(
        [_linha_csv(telefone="600000101", email="", classe_site="proprio")], **_kwargs_comuns(tmp_path)
    )
    linha = linhas[0]
    assert linha["Mensagem sugerida"] == ""
    assert linha["WhatsApp"] != ""
    assert "text=" not in linha["WhatsApp"]


def test_linha_geral_canal_email_preenche_abrir_email_nao_whatsapp(tmp_path):
    linhas = montar_linhas_geral([_linha_csv(telefone="", email="x@y.es")], **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Canal"] == "email"
    assert linha["Abrir e-mail"].startswith("mailto:x@y.es")
    assert linha["WhatsApp"] == ""


def test_linha_geral_canal_para_depois_sem_nenhum_link(tmp_path):
    linhas = montar_linhas_geral([_linha_csv(telefone="", email="")], **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Canal"] == "para depois"
    assert linha["WhatsApp"] == ""
    assert linha["Abrir e-mail"] == ""


def test_linha_geral_profissional_detectado_e_captura_e_contexto(tmp_path):
    (tmp_path / "p1_mobile.png").write_bytes(b"x")
    linhas = montar_linhas_geral(
        [_linha_csv(nome="Dr. Juan Pérez", place_id="p1", classe_site="proprio")], **_kwargs_comuns(tmp_path)
    )
    linha = linhas[0]
    assert linha["Profissional detectado"] == "sim"
    assert linha["Captura"].endswith("p1_mobile.png")
    assert "ÂNGULO: sem_angulo" in linha["Contexto para IA"]  # "proprio" não tem modelo na Direta neste MVP


# --- montar_linhas_nata: ângulo contato / poucas_avaliacoes / lentidao ------


def test_linha_nata_angulo_contato(tmp_path):
    lead = _lead_nata(telefone_na_pagina="texto")
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Ângulo"] == "contato"
    assert linha["mensagem"] != ""
    assert linha["Status"] == ""
    assert linha["pista"] == "nata"


def test_linha_nata_whatsapp_com_mensagem_valida_leva_o_texto(tmp_path):
    """Mesma coluna única "WhatsApp" da aba Geral -- com celular e mensagem,
    o link já vem com `?text=`."""
    lead = _lead_nata(telefone_na_pagina="texto")
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["mensagem"] != ""
    assert "text=" in linha["WhatsApp"]


def test_linha_nata_whatsapp_sem_angulo_usa_conversa_vazia(tmp_path):
    """Controle negativo: sem ângulo (sem mensagem), mas com celular --
    "WhatsApp" ainda é o link de conversa vazia, nunca célula em branco."""
    lead = _lead_nata()  # nenhuma regra de ângulo bate
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["mensagem"] == ""
    assert linha["WhatsApp"] != ""
    assert "text=" not in linha["WhatsApp"]


def test_linha_nata_angulo_poucas_avaliacoes_concorrente_em_candidatos_triagem(tmp_path):
    """Prova central: o concorrente que qualifica o ângulo poucas_avaliacoes
    pode estar em candidatos_triagem, não só em nata (limiar 10x)."""
    lead = _lead_nata(place_id="p1", nota_estado="CONFIRMADO_PRESENTE", nota_valor=4.8,
                      avaliacoes_estado="CONFIRMADO_PRESENTE", avaliacoes_valor=10)
    concorrente = _lead_nata(place_id="p2", nota_estado="CONFIRMADO_PRESENTE", nota_valor=4.0,
                             avaliacoes_estado="CONFIRMADO_PRESENTE", avaliacoes_valor=110)
    linhas = montar_linhas_nata(
        [lead], [concorrente], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path)
    )
    linha_lead = next(l for l in linhas if l["nome"] == lead["identidade"]["nome"] and l["pista"] == "nata")
    assert linha_lead["Ângulo"] == "poucas_avaliacoes"
    # décima rodada (27/09/2026): a mensagem não cita mais o número do
    # concorrente; o único concorrente (em candidatos_triagem) escolhe o
    # modelo _singular ("con otra")
    assert "Puede encontrarse con otra." in linha_lead["mensagem"]
    assert "110" not in linha_lead["mensagem"]


def test_linha_nata_nome_generico_de_nicho_e_cidade_nao_e_falso_positivo(tmp_path):
    """Prova central do defeito real, 25/09/2026: lead "Psicólogo en
    Alicante - Pedro Ficticio" (nome limpo "Psicólogo en Alicante", só
    palavras genéricas) não pode ter a mensagem rejeitada só porque o
    modelo de poucas_avaliacoes legitimamente cita "quien busca psicólogo
    en Alicante" -- usa config/cidade_da_busca.json real (Alicante)."""
    lead = _lead_nata(
        place_id="p1", nome="Psicólogo en Alicante - Pedro Ficticio", nicho="Psicólogo",
        nota_estado="CONFIRMADO_PRESENTE", nota_valor=4.8,
        avaliacoes_estado="CONFIRMADO_PRESENTE", avaliacoes_valor=10,
    )
    concorrente = _lead_nata(place_id="p2", nicho="Psicólogo", nota_estado="CONFIRMADO_PRESENTE", nota_valor=4.0,
                             avaliacoes_estado="CONFIRMADO_PRESENTE", avaliacoes_valor=110)
    linhas = montar_linhas_nata(
        [lead], [concorrente], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path)
    )
    linha_lead = next(l for l in linhas if l["place_id"] == "p1")
    assert linha_lead["Ângulo"] == "poucas_avaliacoes"
    assert linha_lead["mensagem"] != ""
    assert "nome do negócio" not in linha_lead["Aviso"]


def test_linha_nata_angulo_poucas_avaliacoes_concorrente_em_descartados_nao_qualifica():
    """`descartado` não tem bloco `reputacao` no contrato -- nunca
    contribui como concorrente, mesmo estando no lote."""
    lead = _lead_nata(place_id="p1", nicho="Fisioterapia", nota_estado="CONFIRMADO_PRESENTE", nota_valor=4.8,
                      avaliacoes_estado="CONFIRMADO_PRESENTE", avaliacoes_valor=10)
    descartado = {
        "place_id": "p2", "identidade": {"nome": "Concorrente", "nicho": "Fisioterapia", "cidade": None},
        "motivo": "sem_canal_de_contato", "campos_do_corte": {},
    }
    angulo = am.escolher_angulo_nata(lead, [lead, descartado], _REGRAS_ANGULO)
    assert angulo != "poucas_avaliacoes"


def test_linha_nata_angulo_lentidao(tmp_path):
    lead = _lead_nata(lcp_ms=12000, psi_estado="CONFIRMADO_PRESENTE", rodadas=2)
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Ângulo"] == "lentidao"
    assert "segundos" in linha["mensagem"]


def test_linha_nata_sem_angulo_fica_com_status_sem_mensagem(tmp_path):
    lead = _lead_nata()
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    linha = linhas[0]
    assert linha["Ângulo"] == am.SEM_ANGULO
    assert linha["mensagem"] == ""
    assert linha["Status"] == STATUS_SEM_MENSAGEM


def test_linha_nata_conferir_antes_de_enviar_sim_com_fora_do_ar(tmp_path):
    lead = _lead_nata(fora_do_ar=True)
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    assert linhas[0]["conferir_antes_de_enviar"] == "sim"


def test_linha_nata_conferir_antes_de_enviar_nao_sem_fora_do_ar(tmp_path):
    lead = _lead_nata(fora_do_ar=False)
    linhas = montar_linhas_nata([lead], [], [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    assert linhas[0]["conferir_antes_de_enviar"] == "não"


def test_montar_linhas_nata_pista_distingue_nata_de_candidatos_triagem(tmp_path):
    nata = [_lead_nata(place_id="p1")]
    candidatos = [_lead_nata(place_id="p2")]
    linhas = montar_linhas_nata(nata, candidatos, [], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    por_pista = {l["pista"] for l in linhas}
    assert por_pista == {"nata", "candidatos_triagem"}
    assert len(linhas) == 2


def test_montar_linhas_nata_descartados_nunca_viram_linha(tmp_path):
    nata = [_lead_nata(place_id="p1")]
    descartado = {"place_id": "p2", "identidade": {"nome": "X", "nicho": None, "cidade": None}, "motivo": "m", "campos_do_corte": {}}
    linhas = montar_linhas_nata(nata, [], [descartado], regras_angulo=_REGRAS_ANGULO, **_kwargs_comuns(tmp_path))
    assert len(linhas) == 1


# --- gerar_planilha: nunca sobrescreve, exclui já abordados, 3 abas --------


def _csv_humano(tmp_path, linhas):
    import csv

    caminho = tmp_path / "qualificador_20260924-120000.csv"
    with caminho.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(linhas[0].keys()), delimiter=";")
        w.writeheader()
        for linha in linhas:
            w.writerow(linha)
    return caminho


def _kwargs_gerar_planilha(tmp_path, **over):
    base = dict(
        apresentacao=_APRESENTACAO, config_validacao=_CONFIG_VALIDACAO, assuntos_email=_ASSUNTOS_EMAIL,
        opcoes_funil=_OPCOES_FUNIL, regras_angulo=_REGRAS_ANGULO, mensagens_angulo=_MENSAGENS_ANGULO,
        diretorio_capturas=tmp_path / "_capturas_inexistente",
        # Isola a entrega consolidada num diretório de teste -- sem isto,
        # gerar_planilha cairia no COMERCIAL_PLANILHAS_DIR real (comercial/.env)
        # e escreveria na pasta de verdade do diretor durante o teste (mesmo
        # defeito já documentado em env_loader.py pra ANTHROPIC_API_KEY).
        consolidada_dir=tmp_path / "_consolidada_teste",
    )
    base.update(over)
    return base


def test_gerar_planilha_grava_geral_nata_e_como_usar_e_nunca_sobrescreve(tmp_path, monkeypatch):
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv()])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso(nata=[_lead_nata()]))

    saida_dir = tmp_path / "_planilhas"
    agora = datetime(2026, 9, 24, 12, 0, 0, tzinfo=timezone.utc)

    resultado = gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=saida_dir, agora=agora, **_kwargs_gerar_planilha(tmp_path),
    )

    caminho = resultado["caminho"]
    assert caminho.name == "planilha_envio_20260924-120000.xlsx"
    assert caminho.is_file()

    wb = load_workbook(caminho)
    assert wb.sheetnames == ["Geral", "Nata", "Como usar"]
    assert wb["Geral"].cell(row=1, column=1).font.bold is True
    assert wb["Geral"].freeze_panes == "A2"
    assert wb["Nata"].cell(row=2, column=2).value == "Clínica Ejemplo"  # coluna 2 = "nome"
    assert wb["Como usar"].cell(row=2, column=1).value == "Lentidão"

    with pytest.raises(PlanilhaJaExisteError):
        gerar_planilha(
            caminho_csv_humano=caminho_csv, saida_dir=saida_dir, agora=agora, **_kwargs_gerar_planilha(tmp_path),
        )


def test_como_usar_nao_menciona_mais_o_teste_ab_e_classifica_a_resposta(tmp_path, monkeypatch):
    """D12 Fase 5.1 (29/09/2026): a aba "Como usar" deixa de instruir Douglas
    a comparar variantes -- não existe mais experimento de abertura. A
    classificação da resposta (positiva/neutra/negativa) continua, sem a
    ressalva "decida antes de olhar a variante" (não há mais variante)."""
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv()])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso(nata=[_lead_nata()]))

    resultado = gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=tmp_path / "_planilhas",
        agora=datetime(2026, 9, 25, 11, 0, 0, tzinfo=timezone.utc), **_kwargs_gerar_planilha(tmp_path),
    )

    wb = load_workbook(resultado["caminho"])
    titulos = [c.value for c in wb["Como usar"]["A"]]
    textos = [c.value for c in wb["Como usar"]["B"]]
    assert "Teste A/B da abertura" not in titulos
    assert "Comparar A e B" not in titulos
    assert "Classificar a resposta" in titulos
    assert any("resposta neutra" in (t or "") for t in textos)
    assert not any("variante" in (t or "").lower() for t in textos)
    assert not any("Google Maps" in (t or "") for t in textos)


def test_gerar_planilha_exclui_quem_ja_foi_abordado_geral_e_nata_e_candidatos(tmp_path, monkeypatch):
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv(place_id="p1"), _linha_csv(place_id="p2")])
    monkeypatch.setattr(
        contrato_loader, "carregar_lote",
        lambda **kwargs: _lote_falso(
            nata=[_lead_nata(place_id="p1"), _lead_nata(place_id="p2")],
            candidatos_triagem=[_lead_nata(place_id="p3")],
        ),
    )

    caminho_registro = tmp_path / "registro_abordagens.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["place_id", "nome"])
    ws.append(["p1", "Clínica Ejemplo"])
    ws.append(["p3", "Clínica Ejemplo"])
    wb.save(caminho_registro)

    resultado = gerar_planilha(
        caminho_csv_humano=caminho_csv, caminho_registro=caminho_registro,
        saida_dir=tmp_path / "_planilhas", agora=datetime(2026, 9, 24, 13, 0, 0, tzinfo=timezone.utc),
        **_kwargs_gerar_planilha(tmp_path),
    )

    assert resultado["excluidos_geral"] == 1
    assert resultado["excluidos_nata"] == 2  # p1 (nata) + p3 (candidatos_triagem)

    wb2 = load_workbook(resultado["caminho"])
    place_ids_geral = [row[COLUNAS_GERAL.index("place_id")] for row in wb2["Geral"].iter_rows(min_row=2, values_only=True)]
    assert "p1" not in place_ids_geral
    assert "p2" in place_ids_geral

    cab_nata = [c.value for c in wb2["Nata"][1]]
    idx_pista = cab_nata.index("pista")
    idx_nome = cab_nata.index("nome")
    linhas_nata = list(wb2["Nata"].iter_rows(min_row=2, values_only=True))
    pistas_presentes = {l[idx_pista] for l in linhas_nata}
    assert pistas_presentes == {"nata"}  # p1 excluído, p3 (candidatos_triagem) excluído -- só p2/nata sobra


def test_place_ids_ja_abordados_arquivo_ausente_devolve_vazio(tmp_path):
    assert place_ids_ja_abordados(tmp_path / "nao-existe.xlsx") == set()


def test_gerar_planilha_aplica_validacao_de_lista_nas_duas_abas(tmp_path, monkeypatch):
    """Defeito relatado pela supervisão, 24/09/2026: a reescrita da planilha
    tinha derrubado a cobertura de teste das listas de escolha (validação de
    dados) -- este teste protege de novo, contra o `gerar_planilha` inteiro,
    não só `_escrever_aba` isolada."""
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv()])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso(nata=[_lead_nata()]))

    resultado = gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=tmp_path / "_planilhas",
        agora=datetime(2026, 9, 24, 17, 0, 0, tzinfo=timezone.utc), **_kwargs_gerar_planilha(tmp_path),
    )

    wb = load_workbook(resultado["caminho"])
    assert len(wb["Geral"].data_validations.dataValidation) >= 4
    assert len(wb["Nata"].data_validations.dataValidation) >= 4


# --- "Contexto para IA" (decisão do diretor, 24/09/2026) -------------------
#
# Montada por PROGRAMA a partir dos dados já coletados/gerados -- nenhuma
# chamada de LLM aqui. Ordem fixa: LEAD, O QUE MEDIMOS, TEXTO DO SITE,
# ÂNGULO, MENSAGEM (mudou na quinta rodada: sem ESTRATÉGIA/FATO USADO/
# TRECHO DO SITE/MODELO, que dependiam da Etapa E2, agora fora do fluxo).


def test_contexto_ia_nata_traz_as_secoes_na_ordem_certa():
    lead = _lead_nata(texto_site="Ofrecemos fisioterapia deportiva.")
    contexto = montar_contexto_ia_nata(lead, "lentidao", "Hola, buenas. ¿se lo envío?")
    linhas = contexto.split("\n")

    assert linhas[0].startswith("LEAD: Clínica Ejemplo")
    assert linhas[1].startswith("O QUE MEDIMOS: ")
    assert linhas[2] == "TEXTO DO SITE: Ofrecemos fisioterapia deportiva."
    assert linhas[3] == "ÂNGULO: lentidao"
    assert linhas[4] == "MENSAGEM: Hola, buenas. ¿se lo envío?"


def test_contexto_ia_nata_texto_do_site_ausente_diz_nao_medido():
    contexto = montar_contexto_ia_nata(_lead_nata(texto_site=None), am.SEM_ANGULO, None)
    assert "TEXTO DO SITE: não medido" in contexto


def test_contexto_ia_nata_sem_angulo_diz_nenhuma_mensagem():
    contexto = montar_contexto_ia_nata(_lead_nata(), am.SEM_ANGULO, None)
    assert "ÂNGULO: sem_angulo" in contexto
    assert "MENSAGEM: (nenhuma -- sem ângulo)" in contexto


def test_contexto_ia_nata_inclui_a_mensagem():
    """Prova central: a mensagem gerada precisa estar no contexto -- não
    pode ser esquecida."""
    contexto = montar_contexto_ia_nata(_lead_nata(), "contato", "Esta é a mensagem que foi enviada de verdade.")
    assert "Esta é a mensagem que foi enviada de verdade." in contexto


def test_contexto_ia_nata_respeita_o_limite_da_celula_truncando_so_o_texto_do_site():
    texto_gigante = "x" * 40000
    lead = _lead_nata(texto_site=texto_gigante)
    contexto = montar_contexto_ia_nata(lead, "contato", "Esta mensagem tem que sobreviver ao corte.")

    assert len(contexto) <= LIMITE_CELULA_XLSX
    assert "truncado" in contexto
    assert "Esta mensagem tem que sobreviver ao corte." in contexto
    assert "MENSAGEM:" in contexto


def test_contexto_ia_nata_texto_curto_nunca_e_truncado():
    contexto = montar_contexto_ia_nata(_lead_nata(texto_site="Texto curto do site."), am.SEM_ANGULO, None)
    assert "truncado" not in contexto
    assert "Texto curto do site." in contexto


def test_contexto_ia_geral_traz_lead_classe_angulo_e_mensagem():
    linha_geral = {
        "nome": "Clínica Ejemplo", "cidade": "Madrid", "telefone": "600000101",
        "site": "https://doctoralia.es/x", "google_maps_url": "https://maps/x",
        "classe_site": "portal", "Ângulo": "portal", "Mensagem sugerida": "Hola, buenas. ¿Le interesa?",
    }
    contexto = montar_contexto_ia_geral(linha_geral)
    assert "LEAD: Clínica Ejemplo" in contexto
    assert "CLASSE DO SITE: portal" in contexto
    assert "ÂNGULO: portal" in contexto
    assert "MENSAGEM SUGERIDA: Hola, buenas. ¿Le interesa?" in contexto


def test_contexto_ia_geral_sem_angulo_mostra_o_valor_sem_angulo():
    linha_geral = {"nome": "X", "cidade": "", "telefone": "", "site": "", "google_maps_url": "", "classe_site": "", "Ângulo": am.SEM_ANGULO}
    contexto = montar_contexto_ia_geral(linha_geral)
    assert f"ÂNGULO: {am.SEM_ANGULO}" in contexto
    assert "MENSAGEM SUGERIDA: (nenhuma)" in contexto


def test_contexto_ia_geral_angulo_ausente_diz_nenhum():
    """Controle: só quando a chave "Ângulo" vem vazia/ausente de fato (não
    o valor "sem_angulo", que é uma string não vazia) o texto cai no
    default "(nenhum)"."""
    linha_geral = {"nome": "X", "cidade": "", "telefone": "", "site": "", "google_maps_url": "", "classe_site": ""}
    contexto = montar_contexto_ia_geral(linha_geral)
    assert "ÂNGULO: (nenhum)" in contexto


# --- Planilha consolidada (entrega do diretor, fora do repo, acrescentada) --
# Decisão do diretor, 01/10/2026: sempre DUAS entregas -- a interna (programa,
# uma planilha nova por rodada, 3 abas) e a consolidada (diretor, um único
# arquivo que cada rodada ABRE e ACRESCENTA, 2 abas, nunca recriado).


def test_atualizar_planilha_consolidada_cria_do_zero_quando_nao_existe(tmp_path):
    destino = tmp_path / "consolidada"
    linhas_geral = [{"place_id": "g1", "nome": "Geral Um"}]
    linhas_nata = [{"place_id": "n1", "nome": "Nata Um"}]

    resultado = pe.atualizar_planilha_consolidada(
        linhas_geral=linhas_geral, linhas_nata=linhas_nata, diretorio=destino,
    )

    assert resultado["pulado"] is False
    assert resultado["novos_geral"] == 1
    assert resultado["novos_nata"] == 1
    caminho = resultado["caminho"]
    assert caminho == destino / pe.NOME_PLANILHA_CONSOLIDADA
    assert caminho.is_file()

    wb = load_workbook(caminho)
    assert wb.sheetnames == ["Geral", "Nata"]  # sem "Como usar"
    assert wb["Geral"].cell(1, 1).value == COLUNAS_GERAL[0]
    assert wb["Geral"].cell(2, COLUNAS_GERAL.index("nome") + 1).value == "Geral Um"
    assert wb["Nata"].cell(2, COLUNAS_NATA.index("nome") + 1).value == "Nata Um"


def test_atualizar_planilha_consolidada_acrescenta_sem_recriar(tmp_path):
    destino = tmp_path / "consolidada"
    pe.atualizar_planilha_consolidada(
        linhas_geral=[{"place_id": "g1", "nome": "Primeira rodada"}],
        linhas_nata=[{"place_id": "n1", "nome": "Primeira rodada nata"}],
        diretorio=destino,
    )

    resultado = pe.atualizar_planilha_consolidada(
        linhas_geral=[{"place_id": "g2", "nome": "Segunda rodada"}],
        linhas_nata=[],
        diretorio=destino,
    )

    assert resultado["novos_geral"] == 1
    assert resultado["novos_nata"] == 0
    wb = load_workbook(resultado["caminho"])
    nomes_geral = [row[COLUNAS_GERAL.index("nome")] for row in wb["Geral"].iter_rows(min_row=2, values_only=True)]
    assert nomes_geral == ["Primeira rodada", "Segunda rodada"]
    # a aba Nata não ganhou linha em branco nem cabeçalho duplicado na 2ª rodada
    assert wb["Nata"].max_row == 2


def test_atualizar_planilha_consolidada_nao_duplica_place_id_ja_presente(tmp_path):
    destino = tmp_path / "consolidada"
    pe.atualizar_planilha_consolidada(
        linhas_geral=[{"place_id": "g1", "nome": "Já estava"}], linhas_nata=[], diretorio=destino,
    )

    resultado = pe.atualizar_planilha_consolidada(
        linhas_geral=[{"place_id": "g1", "nome": "Já estava"}, {"place_id": "g2", "nome": "Novo"}],
        linhas_nata=[],
        diretorio=destino,
    )

    assert resultado["novos_geral"] == 1  # só g2 -- g1 já estava na aba
    wb = load_workbook(resultado["caminho"])
    place_ids = [row[COLUNAS_GERAL.index("place_id")] for row in wb["Geral"].iter_rows(min_row=2, values_only=True)]
    assert place_ids == ["g1", "g2"]


def test_atualizar_planilha_consolidada_pulada_sem_diretorio_configurado(monkeypatch):
    monkeypatch.delenv(pe.ENV_CONSOLIDADA_DIR, raising=False)
    monkeypatch.setattr(pe.env_loader, "carregar_env", lambda *a, **k: [])

    resultado = pe.atualizar_planilha_consolidada(linhas_geral=[], linhas_nata=[])

    assert resultado["pulado"] is True
    assert pe.ENV_CONSOLIDADA_DIR in resultado["motivo"]


def test_gerar_planilha_tambem_atualiza_a_consolidada(tmp_path, monkeypatch):
    import contrato_loader

    caminho_csv = _csv_humano(tmp_path, [_linha_csv()])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso(nata=[_lead_nata()]))
    consolidada_dir = tmp_path / "_consolidada"

    resultado = gerar_planilha(
        caminho_csv_humano=caminho_csv, saida_dir=tmp_path / "_planilhas",
        agora=datetime(2026, 10, 1, 9, 0, 0, tzinfo=timezone.utc),
        **_kwargs_gerar_planilha(tmp_path, consolidada_dir=consolidada_dir),
    )

    consolidada = resultado["consolidada"]
    assert consolidada["pulado"] is False
    assert consolidada["caminho"] == consolidada_dir / pe.NOME_PLANILHA_CONSOLIDADA
    assert consolidada["novos_geral"] == 1
    assert consolidada["novos_nata"] == 1
    wb = load_workbook(consolidada["caminho"])
    assert wb.sheetnames == ["Geral", "Nata"]


# --- Checagem de hyperlink desalinhado (defeito real, 01/10/2026) ----------
# `ws.delete_rows()` desloca o TEXTO das células mas não realinha o objeto
# `hyperlink` do openpyxl -- achado numa limpeza manual da planilha
# consolidada que corrompeu os links de WhatsApp/e-mail (cada linha abria o
# WhatsApp/e-mail de OUTRO lead). `_verificar_links_aba` é a rede que pega
# isso antes de qualquer planilha sair com link errado.


def _ws_com_link(colunas, linhas_dados, texto_wa=None, hyperlink_wa=None):
    """Monta uma aba mínima com uma linha de dado e controla separadamente o
    TEXTO visível e o destino do hyperlink da célula WhatsApp -- é exatamente
    essa divergência que `_verificar_links_aba` precisa detectar."""
    wb = Workbook()
    ws = wb.active
    ws.append(list(colunas))
    for linha in linhas_dados:
        ws.append([linha.get(c, "") for c in colunas])
    if texto_wa is not None:
        idx = colunas.index("WhatsApp") + 1
        cel = ws.cell(2, idx)
        cel.value = texto_wa
        if hyperlink_wa:
            cel.hyperlink = hyperlink_wa
    return ws


def test_verificar_links_aba_sem_problema_quando_tudo_bate():
    colunas = ("telefone", "email", "WhatsApp", "Abrir e-mail")
    ws = _ws_com_link(
        colunas, [{"telefone": "600000101", "email": "", "WhatsApp": "https://wa.me/34600000101"}],
        texto_wa="https://wa.me/34600000101", hyperlink_wa="https://wa.me/34600000101",
    )
    assert _verificar_links_aba("Geral", ws, colunas) == []


def test_verificar_links_aba_sem_hyperlink_nenhum_nao_e_problema():
    """Controle: linha sem link nenhum (canal não aplicável) não gera queixa --
    ausência de link não é o defeito que esta checagem cobre."""
    colunas = ("telefone", "email", "WhatsApp", "Abrir e-mail")
    ws = _ws_com_link(colunas, [{"telefone": "", "email": "", "WhatsApp": ""}])
    assert _verificar_links_aba("Geral", ws, colunas) == []


def test_verificar_links_aba_detecta_link_da_linha_anterior():
    """Reproduz o defeito real: hyperlink da linha aponta pro telefone de
    OUTRA linha (não o da própria linha) -- exatamente o sintoma do relatório
    (`delete_rows` desloca texto mas não o hyperlink). Texto da célula é
    igual ao hyperlink (ambos "errados" do mesmo jeito) pra isolar só o
    sintoma do telefone trocado, sem also acionar a checagem de texto!=link."""
    colunas = ("telefone", "email", "WhatsApp", "Abrir e-mail")
    ws = _ws_com_link(
        colunas, [{"telefone": "600000301", "email": "", "WhatsApp": "https://wa.me/34600000302"}],
        texto_wa="https://wa.me/34600000302", hyperlink_wa="https://wa.me/34600000302",  # telefone de outro lead
    )
    problemas = _verificar_links_aba("Geral", ws, colunas)
    assert len(problemas) == 1
    assert "não corresponde ao telefone" in problemas[0]


def test_verificar_links_aba_detecta_texto_diferente_do_link():
    colunas = ("telefone", "email", "WhatsApp", "Abrir e-mail")
    ws = _ws_com_link(
        colunas, [{"telefone": "600000101", "email": "", "WhatsApp": "https://wa.me/34600000101"}],
        texto_wa="texto velho", hyperlink_wa="https://wa.me/34600000101",
    )
    problemas = _verificar_links_aba("Geral", ws, colunas)
    assert any("difere do link" in p for p in problemas)


def test_verificar_links_aba_detecta_mailto_sem_email_na_linha():
    colunas = ("telefone", "email", "WhatsApp", "Abrir e-mail")
    ws = Workbook().active
    ws.append(list(colunas))
    ws.append(["", "", "", "mailto:alguem@exemplo.com"])
    idx = colunas.index("Abrir e-mail") + 1
    ws.cell(2, idx).hyperlink = "mailto:alguem@exemplo.com"
    problemas = _verificar_links_aba("Geral", ws, colunas)
    assert any("não tem email" in p for p in problemas)


def test_verificar_consistencia_links_le_arquivo_e_agrupa_por_aba(tmp_path):
    caminho = tmp_path / "planilha.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Geral"
    ws1.append(["telefone", "email", "WhatsApp", "Abrir e-mail"])
    ws1.append(["600000101", "", "https://wa.me/34999999999", ""])
    ws1.cell(2, 3).hyperlink = "https://wa.me/34999999999"  # errado de propósito, texto==link
    ws2 = wb.create_sheet("Nata")
    ws2.append(["telefone", "email", "WhatsApp", "Abrir e-mail"])
    ws2.append(["600000102", "", "https://wa.me/34600000102", ""])
    ws2.cell(2, 3).hyperlink = "https://wa.me/34600000102"  # certo
    wb.save(caminho)

    resultado = verificar_consistencia_links(caminho)

    assert len(resultado["Geral"]) == 1
    assert resultado["Nata"] == []


def test_gerar_planilha_levanta_link_inconsistente_e_nao_grava_nada(tmp_path, monkeypatch):
    """Se `_linha_geral`/`_linha_nata` algum dia produzissem um link
    desalinhado, `gerar_planilha` tem que falhar ANTES de gravar -- mesmo
    espírito de `PlaceIdAusenteError`: nunca uma planilha incompleta/errada
    em silêncio. Simulado forçando `_verificar_links_aba` a sempre achar
    problema, já que reproduzir o desalinhamento real exigiria sujar
    `_linha_geral` só pro teste."""
    import contrato_loader
    import planilha_envio as pe_mod

    caminho_csv = _csv_humano(tmp_path, [_linha_csv()])
    monkeypatch.setattr(contrato_loader, "carregar_lote", lambda **kwargs: _lote_falso())
    monkeypatch.setattr(pe_mod, "_verificar_links_aba", lambda *a, **k: ["problema forçado pelo teste"])

    saida_dir = tmp_path / "_planilhas"
    with pytest.raises(LinkInconsistenteError):
        gerar_planilha(
            caminho_csv_humano=caminho_csv, saida_dir=saida_dir,
            agora=datetime(2026, 10, 1, 9, 0, 0, tzinfo=timezone.utc),
            **_kwargs_gerar_planilha(tmp_path),
        )

    assert list(saida_dir.glob("*.xlsx")) == []  # nada foi gravado


def test_atualizar_planilha_consolidada_levanta_e_nao_salva_se_ja_existia_quebrada(tmp_path, monkeypatch):
    """Mesma garantia na consolidada: se a aba já reaberta tiver um link
    desalinhado (ex.: sobrevivente de uma rodada anterior corrompida), a
    atualização falha e o arquivo em disco NÃO é tocado -- em vez de
    acrescentar linhas boas por cima de um arquivo já quebrado."""
    import planilha_envio as pe_mod

    destino = tmp_path / "consolidada"
    pe.atualizar_planilha_consolidada(
        linhas_geral=[{"place_id": "g1", "nome": "Lead Um", "telefone": "600000101", "WhatsApp": "https://wa.me/34600000101"}],
        linhas_nata=[], diretorio=destino,
    )
    antes = (destino / pe.NOME_PLANILHA_CONSOLIDADA).stat().st_mtime

    monkeypatch.setattr(pe_mod, "_verificar_links_aba", lambda *a, **k: ["problema forçado pelo teste"])
    with pytest.raises(LinkInconsistenteError):
        pe.atualizar_planilha_consolidada(
            linhas_geral=[{"place_id": "g2", "nome": "Lead Dois"}], linhas_nata=[], diretorio=destino,
        )

    depois = (destino / pe.NOME_PLANILHA_CONSOLIDADA).stat().st_mtime
    assert antes == depois  # arquivo em disco não mudou


def test_main_verificar_imprime_ok_e_sai_0_quando_sem_problema(tmp_path, capsys):
    caminho = tmp_path / "planilha.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Geral"
    ws.append(["telefone", "WhatsApp"])
    ws.append(["600000101", "https://wa.me/34600000101"])
    ws.cell(2, 2).hyperlink = "https://wa.me/34600000101"
    wb.save(caminho)

    codigo = pe.main(["--verificar", str(caminho)])

    assert codigo == 0
    assert "OK, nenhum link desalinhado." in capsys.readouterr().out


def test_main_verificar_sai_1_e_lista_problema_quando_ha_desalinhamento(tmp_path, capsys):
    caminho = tmp_path / "planilha.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Geral"
    ws.append(["telefone", "WhatsApp"])
    ws.append(["600000101", "https://wa.me/34999999999"])
    ws.cell(2, 2).hyperlink = "https://wa.me/34999999999"  # texto==link, só o telefone não bate
    wb.save(caminho)

    codigo = pe.main(["--verificar", str(caminho)])

    saida = capsys.readouterr().out
    assert codigo == 1
    assert "não corresponde ao telefone" in saida
    assert "TOTAL: 1 problema" in saida


# --- main() (comando `python ferramentas/planilha_envio.py`) --------------


def test_main_chama_gerar_planilha_com_saida_dir_e_imprime_resumo(tmp_path, monkeypatch, capsys):
    chamada = {}

    def _gerar_planilha_falso(*, saida_dir=None):
        chamada["saida_dir"] = saida_dir
        return {
            "caminho": tmp_path / "planilha_envio_20260924-120000.xlsx",
            "excluidos_geral": 2,
            "excluidos_nata": 1,
            "consolidada": {"pulado": True, "motivo": "COMERCIAL_PLANILHAS_DIR não definida -- consolidada pulada."},
        }

    monkeypatch.setattr(pe, "gerar_planilha", _gerar_planilha_falso)

    codigo = pe.main(["--saida-dir", str(tmp_path)])

    assert codigo == 0
    assert chamada["saida_dir"] == tmp_path
    saida = capsys.readouterr().out
    assert "planilha_envio_20260924-120000.xlsx" in saida
    assert "2 excluído" in saida
    assert "1 excluído" in saida
    assert "[Consolidada]" in saida


def test_main_sem_saida_dir_passa_none_para_gerar_planilha(monkeypatch):
    chamada = {}

    def _gerar_planilha_falso(*, saida_dir=None):
        chamada["saida_dir"] = saida_dir
        return {
            "caminho": Path("planilha.xlsx"), "excluidos_geral": 0, "excluidos_nata": 0,
            "consolidada": {"pulado": True, "motivo": "COMERCIAL_PLANILHAS_DIR não definida -- consolidada pulada."},
        }

    monkeypatch.setattr(pe, "gerar_planilha", _gerar_planilha_falso)

    assert pe.main([]) == 0
    assert chamada["saida_dir"] is None


# --- abertura A/B, aposentada em 29/09/2026 (D12 Fase 5.1) -------------------
#
# A coluna "Variante" e o distribuidor A/B (`angulo_mensagem.
# DistribuidorVariante`) foram removidos do código -- os testes que os
# validavam foram removidos com eles, não afrouxados: não existe mais
# alternância nem coluna para testar. `test_colunas_geral_tem_place_id` e
# afins continuam cobrindo a forma das colunas que sobraram.


def test_colunas_geral_e_nata_nao_tem_mais_variante():
    assert "Variante" not in COLUNAS_GERAL
    assert "Variante" not in COLUNAS_NATA
